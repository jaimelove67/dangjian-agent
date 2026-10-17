"""上级审批 / 备案报送（T-MEM-11）纯函数、输入契约与导出：不需要数据库与会话。"""

import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.rules.member_stages import STAGE_LABELS, MemberStage, load_stage_rules
from app.schemas.member_upper_report import (
    UpperReportCreate,
    UpperReportRevoke,
    UpperReportSettle,
)
from app.services.member_upper_report_service import (
    DEFAULT_APPROVAL_DAYS,
    KIND_LABELS,
    RESULT_LABELS,
    STATUS_LABELS,
    UPPER_ITEM_MAP,
    UPPER_ITEMS,
    UpperItem,
    approval_deadline,
    can_report,
    days_left,
    is_due_soon,
    is_overdue,
    item_spec,
    load_rules,
    report_materials,
    stage_label,
)
from app.services.roster_io import UPPER_REPORT_HEADERS, build_upper_report_export

TODAY = date(2026, 10, 16)
YESTERDAY = (date.today() - timedelta(days=1)).isoformat()
TOMORROW = (date.today() + timedelta(days=1)).isoformat()


class _FakeRow:
    """最小记录替身：只带 `row_view` 与超期判断要读的字段。"""

    def __init__(self, **kwargs: object) -> None:
        self.kind = "approval"
        self.status = "submitted"
        self.submitted_on = TODAY
        self.upper_org = "TEST 院系党委"
        for key, value in kwargs.items():
            setattr(self, key, value)


def test_items_are_keyed_by_target_stage_with_citable_basis():
    # 事项由"目标阶段"唯一决定：支部大会 / 支委会作出决定后才报送，阶段就是天然的前提条件
    assert [item.key for item in UPPER_ITEMS] == [
        "activist_record",
        "candidate_record",
        "probationary_approval",
        "member_approval",
    ]
    assert [item.stage for item in UPPER_ITEMS] == [
        MemberStage.ACTIVIST.value,
        MemberStage.CANDIDATE.value,
        MemberStage.PROBATIONARY.value,
        MemberStage.MEMBER.value,
    ]
    # 备案两项、审批两项：细则第八 / 十三条为备案，第二十一 / 二十二条与第五章为审批
    assert [item.kind for item in UPPER_ITEMS] == ["record", "record", "approval", "approval"]
    # 依据必须可引用、可核对：空话或"系统规定"都不算依据
    for item in UPPER_ITEMS:
        assert "细则" in item.basis
        assert item.label and item.report_to
    assert item_spec("member_approval").kind == "approval"
    # 事项不在目录内要报 422，不能让请求自己编一个事项名进台账
    with pytest.raises(Exception) as excinfo:
        item_spec("custom_approval")
    assert getattr(excinfo.value, "status_code", None) == 422


def test_can_report_requires_stage_reached():
    activist = item_spec("activist_record")
    candidate = item_spec("candidate_record")
    probationary = item_spec("probationary_approval")
    member = item_spec("member_approval")
    # 还没到那个环节就不给办：入党申请人不能报任何备案 / 审批
    for spec in (activist, candidate, probationary, member):
        assert can_report(MemberStage.APPLICANT.value, spec) is False
    # 积极分子只能报"确定积极分子备案"；发展对象可回顾地补报积极分子备案
    assert can_report(MemberStage.ACTIVIST.value, activist) is True
    assert can_report(MemberStage.ACTIVIST.value, candidate) is False
    assert can_report(MemberStage.CANDIDATE.value, activist) is True
    assert can_report(MemberStage.CANDIDATE.value, candidate) is True
    assert can_report(MemberStage.CANDIDATE.value, probationary) is False
    # 正式党员可以补登历史备案 / 审批（存量人员导入后按事实回填）
    assert can_report(MemberStage.MEMBER.value, activist) is True
    assert can_report(MemberStage.MEMBER.value, member) is True
    # 非发展阶段（退回）与非法值都不受理
    assert can_report(MemberStage.REJECTED.value, activist) is False
    assert can_report("legacy-stage", activist) is False


def test_report_materials_union_matches_stage_requirements():
    rules = load_stage_rules()
    config = load_rules()
    # 入党积极分子备案：阶段材料要求（申请人 → 积极分子）之外另加"入党积极分子备案报告"
    activist = item_spec("activist_record")
    materials = report_materials(
        activist, MemberStage.ACTIVIST, rules, config["extra_materials"][activist.key]
    )
    assert materials == ("入党申请书", "党组织谈话记录", "入党积极分子备案报告")
    # 接收预备党员审批：沿发展链累加到"预备党员"为止，再补入党志愿书与培养考察登记表
    probationary = item_spec("probationary_approval")
    materials = report_materials(
        probationary,
        MemberStage.PROBATIONARY,
        rules,
        config["extra_materials"][probationary.key],
    )
    assert list(materials[:4]) == ["入党申请书", "党组织谈话记录", "思想汇报", "培养考察记录"]
    assert "政治审查材料" in materials
    assert "中国共产党入党志愿书" in materials
    assert "入党积极分子培养考察登记表" in materials
    # 去重保序：配置里重复登记同一份材料不会让清单出现两遍
    materials = report_materials(
        probationary, MemberStage.PROBATIONARY, rules, ("政治审查材料", "新加材料")
    )
    assert list(materials).count("政治审查材料") == 1
    assert materials[-1] == "新加材料"
    # 清单与阶段流转 / 组织关系转入同一来源：预备党员档案清单是它的子集
    assert set(rules.materials_for_stage(MemberStage.PROBATIONARY)) <= set(materials)


def test_rules_come_from_config_and_are_adjustable():
    rules = load_rules()
    # 审批时限取自细则第二十五条（三个月内审批），配置可改；备忘：细则未统一规定备案时限
    assert rules["approval_days"] == 90
    assert rules["remind_advance_days"] == 15
    assert DEFAULT_APPROVAL_DAYS == 90
    assert set(rules["extra_materials"]) == {
        "activist_record",
        "candidate_record",
        "probationary_approval",
        "member_approval",
    }


def test_overdue_only_applies_to_pending_approvals():
    # 审批类：超过配置的批复时限仍未回填才算超期，到期当天不算
    row = _FakeRow(kind="approval", status="submitted", submitted_on=TODAY - timedelta(days=90))
    assert is_overdue(row, 90, TODAY) is False
    assert days_left(row, 90, TODAY) == 0
    row.submitted_on = TODAY - timedelta(days=91)
    assert is_overdue(row, 90, TODAY) is True
    assert days_left(row, 90, TODAY) == -1
    # 已回填、已作废都不再提示：办结的事不该一直挂着红
    assert is_overdue(_FakeRow(status="settled", submitted_on=TODAY), 90, TODAY) is False
    assert is_overdue(_FakeRow(status="revoked", submitted_on=TODAY), 90, TODAY) is False
    # 备案类细则未统一规定时限：不计到期、也不显示倒计时，避免编一个不存在的期限
    record = _FakeRow(kind="record", submitted_on=TODAY - timedelta(days=365))
    assert is_overdue(record, 90, TODAY) is False
    assert days_left(record, 90, TODAY) is None
    assert approval_deadline(date(2026, 1, 1), 90) == date(2026, 4, 1)


def test_due_soon_warns_before_the_deadline_passes():
    # 到期前提前提示：不能等到超期了才第一次提醒（配置的提前天数是真起作用的）
    row = _FakeRow(kind="approval", status="submitted", submitted_on=TODAY - timedelta(days=75))
    assert days_left(row, 90, TODAY) == 15
    assert is_due_soon(row, 90, 15, TODAY) is True
    row.submitted_on = TODAY - timedelta(days=74)
    assert days_left(row, 90, TODAY) == 16
    assert is_due_soon(row, 90, 15, TODAY) is False
    # 已经超期的走 overdue，不再算"临近"
    row.submitted_on = TODAY - timedelta(days=400)
    assert is_due_soon(row, 90, 15, TODAY) is False
    # 提前天数配置为 0 时不提示：只剩 10 天也只在配置了提前量时才提醒
    soon = _FakeRow(kind="approval", status="submitted", submitted_on=TODAY - timedelta(days=80))
    assert days_left(soon, 90, TODAY) == 10
    assert is_due_soon(soon, 90, 15, TODAY) is True
    assert is_due_soon(soon, 90, 0, TODAY) is False
    # 已回填 / 已作废不再提示；备案类没有时限，也就没有"临近"
    assert is_due_soon(_FakeRow(status="settled", submitted_on=TODAY), 90, 15, TODAY) is False
    assert is_due_soon(_FakeRow(kind="record", submitted_on=TODAY), 90, 15, TODAY) is False


def test_labels_and_stage_snapshot_are_readable():
    assert KIND_LABELS == {"record": "备案", "approval": "审批"}
    assert STATUS_LABELS["submitted"] == "已报送待批复"
    assert STATUS_LABELS["settled"] == "已回填批复"
    assert STATUS_LABELS["revoked"] == "已作废"
    assert RESULT_LABELS["agreed"] == "同意 / 批准"
    assert RESULT_LABELS["supplement"] == "需补充材料"
    assert stage_label("probationary") == STAGE_LABELS[MemberStage.PROBATIONARY]
    # 历史行的阶段值可能已不在枚举内：读不出来留空，而不是抛错让台账打不开
    assert stage_label("legacy-stage") == ""
    assert stage_label(None) == ""


def test_settle_contract_requires_a_citable_reply():
    # 没有文号的"同意 / 不同意"不是批复：不能只凭一次点击把上级结论写成"同意"
    with pytest.raises(ValidationError):
        UpperReportSettle(result="agreed", reply_no="", reason="收到口头通知")
    with pytest.raises(ValidationError):
        UpperReportSettle(result="rejected", reply_no="", reason="收到电话")
    # "需补充材料"必须写清缺什么，否则回填就是走过场
    with pytest.raises(ValidationError):
        UpperReportSettle(result="supplement", decide_note="", reason="先记一下")
    ok = UpperReportSettle(
        result="agreed", reply_no="校党组〔2026〕15号", decided_on=YESTERDAY, reason="收到批复"
    )
    assert ok.reply_no == "校党组〔2026〕15号"
    assert UpperReportSettle(
        result="supplement", decide_note="缺政治审查材料的调查结论", reason="需补件"
    ).decide_note
    # 批复日期不能晚于今天
    with pytest.raises(ValidationError):
        UpperReportSettle(result="agreed", reply_no="×号", decided_on=TOMORROW, reason="先记")
    # 请求里没有、也不该有"批准"字段
    with pytest.raises(ValidationError):
        UpperReportSettle(result="approved", reply_no="×号", reason="先记")


def test_create_and_revoke_contracts_strip_and_reject_invalid_input():
    body = UpperReportCreate(
        profile_id="p1",
        item="candidate_record",
        upper_org="  TEST 院系党委  ",
        letter_no=" 报字2026-01 ",
        submitted_on=YESTERDAY,
        reason="  报送备案  ",
    )
    # 前后空格统一去掉：带上空格的文号会让"对账"变成两回事
    assert body.upper_org == "TEST 院系党委"
    assert body.letter_no == "报字2026-01"
    assert body.reason == "报送备案"
    with pytest.raises(ValidationError):
        UpperReportCreate(
            profile_id="p1",
            item="candidate_record",
            upper_org="党委",
            submitted_on=TOMORROW,
            reason="提前报",
        )
    with pytest.raises(ValidationError):
        UpperReportCreate(
            profile_id="p1",
            item="not_an_item",
            upper_org="党委",
            submitted_on=YESTERDAY,
            reason="乱填事项",
        )
    with pytest.raises(ValidationError):
        UpperReportCreate(
            profile_id="p1",
            item="candidate_record",
            upper_org="党委",
            submitted_on=YESTERDAY,
            reason="报送",
            approved=True,  # 不存在"由系统批准"这种字段
        )
    with pytest.raises(ValidationError):
        UpperReportRevoke(reason="   ")


def test_export_lists_checklist_not_material_contents():
    rows = [
        {
            "item": "probationary_approval",
            "item_label": item_spec("probationary_approval").label,
            "kind_label": "审批",
            "member_name": "合成周砚青",
            "org_name": "TEST 支部",
            "stage_label": "预备党员",
            "upper_org": "TEST 院系党委",
            "letter_no": "",
            "submitted_on": date(2026, 7, 1),
            "materials": ["入党申请书", "政治审查材料", "中国共产党入党志愿书"],
            "status_label": "已报送待批复",
            "result_label": "",
            "decided_on": None,
            "reply_no": "",
            "decide_note": "",
            "overdue": True,
        }
    ]
    content = build_upper_report_export(rows, scope_note="TEST 支部")
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["报送台账", "导出说明"]
    sheet = book["报送台账"]
    assert tuple(cell.value for cell in sheet[1]) == UPPER_REPORT_HEADERS
    values = [cell.value for cell in sheet[2]]
    assert values[0] == item_spec("probationary_approval").label
    assert values[1] == "审批"
    assert values[7] == "2026-07-01"
    # 随报材料导出的是**清单**（顿号连接），不是材料原文：导出件不会带出政审材料内容
    assert values[8] == "入党申请书、政治审查材料、中国共产党入党志愿书"
    assert values[14] == "是"
    guide = "\n".join(
        str(cell.value) for row in book["导出说明"].iter_rows() for cell in row if cell.value
    )
    assert "不含材料原件" in guide
    assert "不代替上级党组织审批" in guide


def test_export_drops_third_party_fields_but_keeps_reply_note():
    """T-MEM-12：第三人字段不落表；批复说明（decide_note）是台账要素，保留。"""
    content = build_upper_report_export(
        [
            {
                "item": "member_approval",
                "item_label": item_spec("member_approval").label,
                "kind_label": "审批",
                "member_name": "合成周砚青",
                "org_name": "TEST 支部",
                "stage_label": "预备党员",
                "upper_org": "TEST 院系党委",
                "letter_no": "TEST-2026-01",
                "submitted_on": date(2026, 7, 1),
                "materials": ["中国共产党入党志愿书"],
                "status_label": "已批复",
                "result_label": "同意",
                "decided_on": date(2026, 7, 20),
                "reply_no": "TEST-批-2026-01",
                "decide_note": "院系党委已批复",
                "overdue": False,
                "note": "经向其舅舅王某某核实",
                "description": "第三人说明",
                "comment": "投票意见",
                "peer_review": "党员互评原文",
            }
        ]
    )
    book = load_workbook(io.BytesIO(content))
    dumped = "\n".join(
        str(cell.value)
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    for marker in ("王某某", "第三人说明", "投票意见", "党员互评原文"):
        assert marker not in dumped
    assert "院系党委已批复" in dumped


def test_export_handles_missing_dates_without_none():
    content = build_upper_report_export(
        [
            {
                "item": "activist_record",
                "item_label": "确定入党积极分子（报上级党委备案）",
                "kind_label": "备案",
                "member_name": "张三",
                "org_name": "TEST 支部",
                "stage_label": "入党积极分子",
                "upper_org": "TEST 院系党委",
                "letter_no": "",
                "submitted_on": date(2026, 7, 1),
                "materials": [],
                "status_label": "已报送待批复",
                "result_label": "",
                "decided_on": None,
                "reply_no": "",
                "decide_note": "",
                "overdue": False,
            }
        ]
    )
    sheet = load_workbook(io.BytesIO(content))["报送台账"]
    values = [cell.value for cell in sheet[2]]
    # 空批复列导出为空而不是 None / "None"：导出件直接可交接（openpyxl 读回空单元格为 None）
    assert values[11] in ("", None)
    assert values[12] in ("", None)
    assert values[14] in ("", None)


def test_item_map_is_the_single_source():
    assert UPPER_ITEM_MAP["activist_record"] == UPPER_ITEMS[0]
    assert isinstance(UPPER_ITEM_MAP["member_approval"], UpperItem)
