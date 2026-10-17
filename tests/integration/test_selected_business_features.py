"""隔离 PostgreSQL 库验证移植功能的权限、事务、迁移和正式 HTTP 接口。"""

import asyncio
import io
import json
import uuid
from contextlib import asynccontextmanager
from datetime import date, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from alembic import command
from alembic.config import Config
from alembic.script import ScriptDirectory
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.pool import NullPool

from app.core.tenant import tenant_context
from app.db.session import get_db
from app.deps import get_current_tenant, get_current_user
from app.llm.service import ModelService
from app.main import app
from app.models.audit import AuditLog
from app.models.knowledge import EmbeddingChunk, KnowledgeDoc
from app.models.member import MemberProfile
from app.models.org import OrgUnit
from app.models.qa_session import QASession
from app.models.user import User, UserRole
from app.rag.retrieval.access import knowledge_filters
from app.core.constants import XLSX_MEDIA_TYPE
from app.rag.retrieval.base import RetrievalResult
from app.rag.retrieval.hybrid import HybridRetriever
from app.rag.retrieval.keyword import KeywordRetriever
from app.services.roster_io import GUIDE_SHEET, HEADERS, TEMPLATE_SHEET
from tests.integration import test_repaired_storage as storage_fixtures

storage_url = storage_fixtures.storage_url


@pytest.fixture(scope="module")
def business_url(storage_url):
    async def prepare():
        engine = create_async_engine(storage_url, poolclass=NullPool, hide_parameters=True)
        try:
            async with engine.begin() as connection:

                def migrate(sync_connection, revision):
                    config = Config()
                    config.set_main_option("script_location", "migrations")
                    config.attributes["connection"] = sync_connection
                    command.upgrade(config, revision)
                    return ScriptDirectory.from_config(config).get_current_head()

                await connection.run_sync(migrate, "004")
                # 模拟协作者 004 已部署并保存过没有 doc_id/index 的旧引用。
                await connection.execute(
                    text(
                        "INSERT INTO qa_sessions "
                        "(id, tenant_id, user_id, question, answer, citations) "
                        "VALUES ('legacy-qa', 'legacy-tenant', 'legacy-user', "
                        "'旧问题', '保留旧回答', :citations)"
                    ),
                    {
                        "citations": json.dumps(
                            [
                                {
                                    "title": "旧制度",
                                    "issuer": "旧单位",
                                    "content": "旧片段",
                                    "score": 0.9,
                                }
                            ]
                        )
                    },
                )
                expected_revision = await connection.run_sync(migrate, "head")
                assert (
                    await connection.scalar(text("SELECT version_num FROM alembic_version"))
                    == expected_revision
                )
            async with AsyncSession(engine, expire_on_commit=False) as db:
                orgs = [
                    ("school-a", "tenant-a", "school", None, False),
                    ("dept-a", "tenant-a", "department", "school-a", False),
                    ("branch-a", "tenant-a", "branch", "dept-a", False),
                    ("dept-b", "tenant-a", "department", "school-a", False),
                    ("branch-b", "tenant-a", "branch", "dept-b", False),
                    ("deleted-branch", "tenant-a", "branch", "dept-a", True),
                    ("hidden-child", "tenant-a", "branch", "deleted-branch", False),
                    ("foreign", "tenant-b", "branch", "dept-a", False),
                ]
                for org_id, tenant, kind, parent, deleted in orgs:
                    db.add(
                        OrgUnit(
                            id=org_id,
                            tenant_id=tenant,
                            name="组织-" + org_id,
                            org_type=kind,
                            parent_id=parent,
                            path=org_id,
                            is_deleted=deleted,
                        )
                    )
                for record_id, tenant, org, deleted, active in [
                    ("a", "tenant-a", "branch-a", False, True),
                    ("dept", "tenant-a", "dept-a", False, True),
                    ("b", "tenant-a", "branch-b", False, True),
                    ("unbound", "tenant-a", None, False, True),
                    ("deleted", "tenant-a", "branch-a", True, True),
                    ("inactive", "tenant-a", "branch-a", False, False),
                    ("hidden", "tenant-a", "hidden-child", False, True),
                    ("foreign", "tenant-b", "foreign", False, True),
                ]:
                    db.add(
                        MemberProfile(
                            id=record_id,
                            tenant_id=tenant,
                            name="合成人员-" + record_id,
                            org_name="合成组织",
                            org_unit_id=org,
                            current_stage="activist",
                            stage_joined_on=date.today() - timedelta(days=400),
                            is_deleted=deleted,
                            is_active=active,
                        )
                    )
                await db.commit()
        finally:
            await engine.dispose()

    asyncio.run(prepare())
    return storage_url


def user(role=UserRole.BRANCH_SECRETARY, org="branch-a", tenant="tenant-a", user_id=None):
    return User(
        id=user_id or str(uuid.uuid4()),
        tenant_id=tenant,
        username="synthetic-reviewer",
        name="合成用户",
        role=role,
        org_unit_id=org,
        is_active=True,
        is_deleted=False,
    )


@asynccontextmanager
async def client_for(url, current_user):
    engine = create_async_engine(url, poolclass=NullPool, hide_parameters=True)
    try:
        async with AsyncSession(engine, expire_on_commit=False) as db:

            async def database():
                yield db

            app.dependency_overrides[get_db] = database
            app.dependency_overrides[get_current_user] = lambda: current_user
            app.dependency_overrides[get_current_tenant] = lambda: current_user.tenant_id
            with tenant_context(current_user.tenant_id, user_id=current_user.id):
                async with httpx.AsyncClient(
                    transport=httpx.ASGITransport(app=app), base_url="http://test"
                ) as client:
                    yield client, db
    finally:
        app.dependency_overrides.clear()
        await engine.dispose()


@pytest.mark.parametrize(
    "role,org,expected",
    [
        (UserRole.BRANCH_SECRETARY, "branch-a", {"a"}),
        (UserRole.ORGANIZER, "branch-a", {"a"}),
        (UserRole.DEPARTMENT_ADMIN, "dept-a", {"a", "dept"}),
        (UserRole.SCHOOL_ADMIN, "school-a", {"a", "dept", "b"}),
        (UserRole.SYSTEM_ADMIN, None, {"a", "dept", "b", "unbound", "hidden"}),
        (UserRole.BRANCH_SECRETARY, None, set()),
        (UserRole.BRANCH_SECRETARY, "foreign", set()),
        (UserRole.DEPARTMENT_ADMIN, "branch-a", set()),
    ],
)
async def test_roster_respects_role_org_tenant_and_active_flags(business_url, role, org, expected):
    async with client_for(business_url, user(role, org)) as (client, _):
        response = await client.get("/api/v1/member/roster")
        assert response.status_code == 200
        data = response.json()["data"]
        assert {item["id"] for item in data["items"]} == expected
        assert data["total"] == len(expected)
        assert all(item["days_in_stage"] == 400 for item in data["items"])


@pytest.mark.parametrize("role", [UserRole.MEMBER, UserRole.APPLICANT])
async def test_ordinary_user_cannot_read_or_create_roster(business_url, role):
    async with client_for(business_url, user(role)) as (client, _):
        assert (await client.get("/api/v1/member/roster")).status_code == 403
        assert (await client.get("/api/v1/member/org-units")).status_code == 403
        assert (
            await client.post("/api/v1/member/roster", json={"name": "合成人员"})
        ).status_code == 403


async def test_roster_pagination_and_org_selector_use_same_scope(business_url):
    async with client_for(business_url, user(UserRole.DEPARTMENT_ADMIN, "dept-a")) as (client, _):
        first = (await client.get("/api/v1/member/roster?page_size=1")).json()["data"]
        second = (await client.get("/api/v1/member/roster?page_size=1&page=2")).json()["data"]
        assert first["total"] == second["total"] == 2
        assert {first["items"][0]["id"], second["items"][0]["id"]} == {"a", "dept"}
        orgs = (await client.get("/api/v1/member/org-units")).json()["data"]
        assert {org["id"] for org in orgs} == {"dept-a", "branch-a"}
        assert (await client.get("/api/v1/member/roster?page=0")).status_code == 422


async def test_roster_filters_by_one_org_within_scope(business_url):
    """T-MEM-02：能看到多个支部的账号，名册要能只按某一个支部看。

    列表、导出、名册统计三处必须用**同一组**筛选（同一权限范围内）；否则会出现
    "支部名册 1 人、支部统计 3 人"这种自相矛盾的口径。
    """
    async with client_for(business_url, user(UserRole.SCHOOL_ADMIN, "school-a")) as (client, _):
        everything = (await client.get("/api/v1/member/roster")).json()["data"]
        only_a = (await client.get("/api/v1/member/roster?org_unit_id=branch-a")).json()["data"]
        assert {item["id"] for item in only_a["items"]} == {"a"}
        assert only_a["total"] == 1 < everything["total"]
        stats_all = (await client.get("/api/v1/member-full/stats")).json()["data"]
        stats_a = (await client.get("/api/v1/member-full/stats?org_unit_id=branch-a")).json()["data"]
        assert stats_all["total"] == everything["total"]
        assert stats_a["total"] == only_a["total"] == 1
        assert stats_a["org_unit_id"] == "branch-a"
        # 导出与列表共用同一组筛选条件：范围外的组织同样 403，不生成文件
        assert (
            await client.get("/api/v1/member/roster/export?org_unit_id=foreign")
        ).status_code == 403


@pytest.mark.parametrize("org_id", ["foreign", "deleted-branch", "missing", "unbound"])
@pytest.mark.parametrize("path", ["/api/v1/member/roster", "/api/v1/member-full/stats"])
async def test_roster_org_filter_rejects_out_of_scope_org(business_url, path, org_id):
    """范围外 / 已删除 / 不存在 / 没有组织的 id 都不能当筛选条件：一律 403，不回退成全量。"""
    async with client_for(business_url, user(UserRole.SCHOOL_ADMIN, "school-a")) as (client, _):
        assert (await client.get(path, params={"org_unit_id": org_id})).status_code == 403


@pytest.mark.parametrize("org_id", ["branch-b", "foreign", "deleted-branch", "missing"])
async def test_roster_creation_rejects_out_of_scope_org(business_url, org_id):
    async with client_for(business_url, user()) as (client, db):
        before = await db.scalar(select(func.count(MemberProfile.id)))
        response = await client.post(
            "/api/v1/member/roster", json={"name": "合成人员", "org_unit_id": org_id}
        )
        assert response.status_code == 403
        assert await db.scalar(select(func.count(MemberProfile.id))) == before


@pytest.mark.parametrize(
    "payload",
    [
        {"name": "   "},
        {"name": "合成人员", "current_stage": "unknown"},
        {"name": "合成人员", "stage_joined_on": "2999-01-01"},
        {"name": "合成人员", "pending": 2147483648},
        {"name": "合成人员", "materials": [""]},
    ],
)
async def test_invalid_member_input_is_rejected_before_write(business_url, payload):
    async with client_for(business_url, user()) as (client, _):
        assert (await client.post("/api/v1/member/roster", json=payload)).status_code == 422


async def test_member_registration_is_persisted_canonical_and_audited(business_url):
    async with client_for(business_url, user()) as (client, db):
        response = await client.post(
            "/api/v1/member/roster",
            json={
                "name": " 合成新台账 ",
                "org_name": "伪造其他组织",
                "materials": ["材料甲", "材料甲"],
            },
        )
        assert response.status_code == 200
        item = response.json()["data"]
        assert item["org_name"] == "组织-branch-a" and item["org_unit_id"] == "branch-a"
        assert item["name"] == "合成新台账" and item["days_in_stage"] == 0
        assert item["materials"] == ["材料甲"]
        profile = await db.get(MemberProfile, item["id"])
        assert profile.current_stage == "applicant"
        audit = (
            await db.execute(select(AuditLog).where(AuditLog.resource_id == item["id"]))
        ).scalar_one()
        assert audit.action == "create" and audit.data_level == "sensitive"
        assert "合成新台账" not in str(audit.new_value)
        # 删除本用例新增台账的可见性，不影响其他用例固定名册预期。
        profile.is_deleted = True
        await db.commit()


async def add_document(db, *, org="dept-a", tenant="tenant-a", shared=False):
    identifier = "soft-delete-" + uuid.uuid4().hex
    document = KnowledgeDoc(
        tenant_id=tenant,
        doc_id=identifier,
        file_name="synthetic.txt",
        title="软删除测试制度",
        issuer="合成单位",
        level="central" if shared else "school",
        visibility="public" if shared else "school",
        security_level="public",
        effective_date=date(2020, 1, 1),
        status="effective",
        tags=[],
        doc_metadata={"org_unit_id": org},
    )
    db.add(document)
    await db.flush()
    db.add(
        EmbeddingChunk(
            tenant_id=tenant,
            doc_id=document.id,
            chunk_id=identifier + "-1",
            content="软删除测试专用规则。",
            sequence=1,
        )
    )
    await db.commit()
    return document


async def test_soft_delete_stops_listing_and_retrieval_but_keeps_audit_and_rows(business_url):
    current_user = user(UserRole.DEPARTMENT_ADMIN, "dept-a")
    async with client_for(business_url, current_user) as (client, db):
        document = await add_document(db)
        assert await KeywordRetriever(db).retrieve(
            "软删除测试专用规则",
            filters={**knowledge_filters(current_user), "doc_ids": [document.doc_id]},
        )
        response = await client.delete("/api/v1/knowledge-docs/" + document.doc_id)
        assert response.status_code == 200
        await db.refresh(document)
        assert document.is_deleted and document.status == "abolished"
        chunks = (
            (await db.execute(select(EmbeddingChunk).where(EmbeddingChunk.doc_id == document.id)))
            .scalars()
            .all()
        )
        assert chunks and all(chunk.is_deleted for chunk in chunks)
        assert not await KeywordRetriever(db).retrieve(
            "软删除测试专用规则",
            filters={
                **knowledge_filters(current_user),
                "doc_ids": [document.doc_id],
                "include_expired": True,
            },
        )
        for suffix in ("", "/chunks"):
            assert (
                await client.get("/api/v1/knowledge-docs/" + document.doc_id + suffix)
            ).status_code == 404
        assert (await client.delete("/api/v1/knowledge-docs/" + document.doc_id)).status_code == 404
        listed = (await client.get("/api/v1/knowledge-docs", params={"q": document.doc_id})).json()[
            "data"
        ]
        assert listed["total"] == 0 and listed["items"] == []
        audit = (
            await db.execute(
                select(AuditLog).where(
                    AuditLog.resource_id == document.id, AuditLog.action == "delete"
                )
            )
        ).scalar_one()
        assert audit.new_value["is_deleted"] is True


@pytest.mark.parametrize(
    "org,tenant,shared",
    [
        ("dept-b", "tenant-a", False),
        (None, "tenant-a", False),
        ("foreign", "tenant-b", True),
    ],
)
async def test_delete_does_not_turn_read_access_into_write_access(
    business_url, org, tenant, shared
):
    async with client_for(business_url, user(UserRole.DEPARTMENT_ADMIN, "dept-a")) as (client, db):
        document = await add_document(db, org=org, tenant=tenant, shared=shared)
        assert (await client.delete("/api/v1/knowledge-docs/" + document.doc_id)).status_code == 404
        await db.refresh(document)
        assert document.is_deleted is False


async def test_ordinary_user_cannot_delete_knowledge(business_url):
    async with client_for(business_url, user(UserRole.MEMBER)) as (client, _):
        assert (await client.delete("/api/v1/knowledge-docs/nonexistent")).status_code == 403


async def test_old_collaborator_history_survives_upgrade_and_remains_readable(business_url):
    async with client_for(
        business_url, user(UserRole.MEMBER, None, "legacy-tenant", "legacy-user")
    ) as (client, _):
        data = (await client.get("/api/v1/qa/sessions")).json()["data"]
        assert data["total"] == 1
        record = data["items"][0]
        assert record["answer"] == "保留旧回答" and record["refused"] is True
        assert record["citations"][0]["index"] == 1
        assert record["citations"][0]["doc_id"] == ""
        assert record["warnings"] and record["disclaimer"]


async def test_persistent_history_is_private_and_paginated(business_url):
    current_user = user(UserRole.MEMBER, None)
    async with client_for(business_url, current_user) as (client, db):
        for tenant, owner, deleted in [
            ("tenant-a", current_user.id, False),
            ("tenant-a", current_user.id, False),
            ("tenant-a", "other-owner", False),
            ("tenant-b", current_user.id, False),
            ("tenant-a", current_user.id, True),
        ]:
            db.add(
                QASession(
                    tenant_id=tenant,
                    user_id=owner,
                    question="合成问题",
                    answer="合成答案",
                    is_deleted=deleted,
                )
            )
        await db.commit()
        first = (await client.get("/api/v1/qa/sessions?page_size=1")).json()["data"]
        second = (await client.get("/api/v1/qa/sessions?page_size=1&page=2")).json()["data"]
        assert first["total"] == second["total"] == 2
        assert first["items"][0]["id"] != second["items"][0]["id"]
        assert (await client.get("/api/v1/qa/sessions?page_size=101")).status_code == 422


async def test_owner_can_soft_delete_own_history_only(business_url):
    """本人可删除自己的记录；他人记录、跨租户与不存在的记录一律 404 且不误删。"""
    current_user = user(UserRole.MEMBER, None)
    async with client_for(business_url, current_user) as (client, db):
        db.add_all(
            [
                QASession(
                    id="own-record",
                    tenant_id="tenant-a",
                    user_id=current_user.id,
                    question="合成问题",
                    answer="合成答案",
                ),
                QASession(
                    id="other-record",
                    tenant_id="tenant-a",
                    user_id="other-owner",
                    question="他人问题",
                    answer="他人答案",
                ),
                QASession(
                    id="cross-record",
                    tenant_id="tenant-b",
                    user_id=current_user.id,
                    question="跨租户",
                    answer="跨租户答案",
                ),
            ]
        )
        await db.commit()

        assert (await client.delete("/api/v1/qa/sessions/does-not-exist")).status_code == 404
        assert (await client.delete("/api/v1/qa/sessions/other-record")).status_code == 404
        assert (await client.delete("/api/v1/qa/sessions/cross-record")).status_code == 404

        assert (await client.delete("/api/v1/qa/sessions/own-record")).status_code == 200
        assert (await client.delete("/api/v1/qa/sessions/own-record")).status_code == 404

        listed = (await client.get("/api/v1/qa/sessions")).json()["data"]
        assert "own-record" not in {item["id"] for item in listed["items"]}

        audit = (
            (
                await db.execute(
                    select(AuditLog).where(
                        AuditLog.user_id == current_user.id,
                        AuditLog.action == "delete",
                        AuditLog.resource_type == "qa_session",
                    )
                )
            )
            .scalars()
            .all()
        )
        assert [entry.resource_id for entry in audit] == ["own-record"]

    # 他人记录未被误删：换到该记录所属账号查看
    async with client_for(
        business_url, user(UserRole.MEMBER, None, "tenant-a", "other-owner")
    ) as (other_client, _):
        data = (await other_client.get("/api/v1/qa/sessions")).json()["data"]
        assert "other-record" in {item["id"] for item in data["items"]}

    # 跨租户记录未被误删：换到该记录所属租户查看
    async with client_for(
        business_url, user(UserRole.MEMBER, None, "tenant-b", current_user.id)
    ) as (cross_client, _):
        data = (await cross_client.get("/api/v1/qa/sessions")).json()["data"]
        assert "cross-record" in {item["id"] for item in data["items"]}


@pytest.mark.parametrize("fail_history", [False, True])
async def test_real_qa_preserves_complete_history_or_recovers_from_failed_flush(
    business_url, monkeypatch, fail_history, caplog
):
    current_user = user(UserRole.MEMBER, None)
    source = RetrievalResult(
        "qa-chunk",
        "第一条 要求提交申请。",
        0.9,
        "qa-document",
        article="第一条",
        metadata={
            "title": "合成制度",
            "issuer": "合成单位",
            "file_name": "synthetic.txt",
            "security_level": "internal",
            "effective_date": "2020-01-01",
            "expiration_date": (date.today() + timedelta(days=10)).isoformat(),
        },
    )
    monkeypatch.setattr(HybridRetriever, "retrieve", AsyncMock(return_value=[source]))
    monkeypatch.setattr(
        ModelService,
        "generate",
        AsyncMock(return_value=SimpleNamespace(content="第一条要求提交申请。[1]")),
    )
    monkeypatch.setattr(
        "app.api.v1.qa._session_store",
        lambda _: SimpleNamespace(get_history=AsyncMock(return_value=[]), append_turn=AsyncMock()),
    )
    async with client_for(business_url, current_user) as (client, db):
        if fail_history:
            original_flush = db.flush

            async def invalid_history_flush(*args, **kwargs):
                for record in db.new:
                    if isinstance(record, QASession):
                        record.question = None  # 真实 NOT NULL 错误，验证 SAVEPOINT 恢复事务。
                await original_flush(*args, **kwargs)

            monkeypatch.setattr(db, "flush", invalid_history_flush)
        response = await client.post(
            "/api/v1/qa",
            json={"question": "合成制度需要哪些材料？", "session_id": "persisted-session"},
        )
        assert response.status_code == 200
        data = response.json()["data"]
        assert data["has_sufficient_evidence"] and data["citations"][0]["doc_id"] == "qa-document"
        history = (await client.get("/api/v1/qa/sessions")).json()["data"]
        audits = (
            (
                await db.execute(
                    select(AuditLog).where(
                        AuditLog.user_id == current_user.id, AuditLog.action == "qa_ask"
                    )
                )
            )
            .scalars()
            .all()
        )
        assert len(audits) == 1
        if fail_history:
            assert history["total"] == 0
            assert "问答历史保存失败，本次回答仍可使用。" in data["warnings"]
            assert "合成制度需要哪些材料" not in caplog.text
            assert "INSERT INTO qa_sessions" not in caplog.text
        else:
            assert history["total"] == 1
            stored = history["items"][0]
            for key, value in data.items():
                assert stored[key] == value
            assert stored["session_id"] == "persisted-session"


# ---------------------------------------------------------------- T-DATA-06 名册导入 / 导出


def roster_xlsx(rows) -> bytes:
    """按模板列顺序造一份名册 Excel（导入用例用）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def roster_file(content: bytes, file_name: str = "roster.xlsx") -> dict:
    return {"file": (file_name, content, XLSX_MEDIA_TYPE)}


IMPORT_FORM = {"source_note": "12371 导出 2026-10-16"}


def guide_text(workbook) -> str:
    return "\n".join(
        str(cell.value) for row in workbook[GUIDE_SHEET].iter_rows() for cell in row if cell.value
    )


async def test_roster_template_and_export_are_readable_xlsx(business_url):
    async with client_for(business_url, user()) as (client, db):
        template = await client.get("/api/v1/member/roster/template")
        assert template.status_code == 200
        assert template.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
        assert "roster-template.xlsx" in template.headers["content-disposition"]
        workbook = load_workbook(io.BytesIO(template.content))
        assert [cell.value for cell in workbook[TEMPLATE_SHEET][1]] == list(HEADERS)
        # 说明页只列当前账号权限范围内的组织，避免填了却导不进来
        guide = guide_text(workbook)
        assert "组织-branch-a" in guide and "组织-dept-b" not in guide

        exported = await client.get("/api/v1/member/roster/export")
        assert exported.status_code == 200
        assert "roster-export.xlsx" in exported.headers["content-disposition"]
        rows = [
            [cell.value for cell in row]
            for row in load_workbook(io.BytesIO(exported.content))[TEMPLATE_SHEET].iter_rows(min_row=2)
        ]
        assert [row[0] for row in rows] == ["合成人员-a"]
        assert rows[0][1] == "合成组织" and rows[0][2] == "入党积极分子"
        audits = (
            (await db.execute(select(AuditLog).where(AuditLog.action == "export_roster"))).scalars().all()
        )
        assert len(audits) == 1 and audits[0].data_level == "sensitive"


@pytest.mark.parametrize("role", [UserRole.MEMBER, UserRole.APPLICANT])
async def test_roster_io_endpoints_require_manage_permission(business_url, role):
    async with client_for(business_url, user(role)) as (client, _):
        assert (await client.get("/api/v1/member/roster/template")).status_code == 403
        assert (await client.get("/api/v1/member/roster/export")).status_code == 403
        assert (
            await client.post(
                "/api/v1/member/roster/import",
                data=IMPORT_FORM,
                files=roster_file(roster_xlsx([])),
            )
        ).status_code == 403


async def test_roster_import_previews_then_commits_with_per_row_report(business_url):
    content = roster_xlsx(
        [
            ["名册导入-甲", "组织-branch-a", "入党积极分子", "2025-09-01", "2025", "2025-2", "1"],
            ["名册导入-乙", "组织-dept-b", "入党积极分子", "2025-09-01", "", "", ""],
            ["名册导入-甲", "组织-branch-a", "正式党员", "2025-09-01", "", "", ""],
            ["合成人员-a", "组织-branch-a", "发展对象", "2025-09-01", "", "", ""],
        ]
    )
    async with client_for(business_url, user()) as (client, db):
        before = await db.scalar(select(func.count(MemberProfile.id)))

        preview = await client.post(
            "/api/v1/member/roster/import",
            data={**IMPORT_FORM, "dry_run": "true"},
            files=roster_file(content),
        )
        assert preview.status_code == 200
        report = preview.json()["data"]
        assert report["dry_run"] is True
        assert report["source_note"] == "12371 导出 2026-10-16"
        assert (report["total_rows"], report["created"], report["skipped"], report["failed"]) == (
            4,
            1,
            1,
            2,
        )
        by_row = {row["row"]: row for row in report["rows"]}
        assert by_row[2]["status"] == "created"
        assert "不在你的权限范围内" in by_row[3]["message"]
        assert "重复" in by_row[4]["message"]
        # 台账里已有的人不因导入改阶段：必须走人工流转
        assert by_row[5]["status"] == "skipped" and "人工流转" in by_row[5]["message"]
        assert "发展对象" in by_row[5]["message"] and "入党积极分子" in by_row[5]["message"]
        # 预检不落库
        assert await db.scalar(select(func.count(MemberProfile.id))) == before

        committed = await client.post(
            "/api/v1/member/roster/import",
            data={**IMPORT_FORM, "dry_run": "false"},
            files=roster_file(content),
        )
        assert committed.status_code == 200
        assert committed.json()["data"]["created"] == 1
        profile = (
            await db.execute(select(MemberProfile).where(MemberProfile.name == "名册导入-甲"))
        ).scalar_one()
        assert profile.org_unit_id == "branch-a" and profile.org_name == "组织-branch-a"
        assert profile.current_stage == "activist" and profile.stage_joined_on == date(2025, 9, 1)
        assert (profile.year, profile.batch_no, profile.pending) == (2025, "2025-2", 1)
        audits = (
            (await db.execute(select(AuditLog).where(AuditLog.action == "import_roster"))).scalars().all()
        )
        assert len(audits) == 2
        assert audits[-1].data_level == "sensitive"
        # 审计按仓库口径只留摘要（业务内容不可读），来源说明因此被绑定进不可篡改的摘要
        assert set(audits[-1].new_value) == {"sha256", "bytes"}

        # 再导一次：已存在的人全部跳过，不会重复建账
        again = await client.post(
            "/api/v1/member/roster/import",
            data={**IMPORT_FORM, "dry_run": "false"},
            files=roster_file(content),
        )
        assert again.status_code == 200
        assert again.json()["data"]["created"] == 0
        assert await db.scalar(
            select(func.count(MemberProfile.id)).where(MemberProfile.name == "名册导入-甲")
        ) == 1

        # 清理：本用例新增的台账不影响同库其他用例的固定名册预期
        profile.is_deleted = True
        await db.commit()


async def test_roster_import_rejects_bad_files_with_actionable_message(business_url):
    async with client_for(business_url, user()) as (client, _):
        short_note = await client.post(
            "/api/v1/member/roster/import",
            data={"source_note": "x"},
            files=roster_file(roster_xlsx([])),
        )
        assert short_note.status_code == 422
        assert "来源说明" in short_note.json()["message"]

        junk = await client.post(
            "/api/v1/member/roster/import",
            data=IMPORT_FORM,
            files=roster_file(b"not an xlsx at all"),
        )
        assert junk.status_code == 422
        assert ".xlsx" in junk.json()["message"] and "另存" in junk.json()["message"]

        headers_only = await client.post(
            "/api/v1/member/roster/import",
            data=IMPORT_FORM,
            files=roster_file(roster_xlsx([])),
        )
        assert headers_only.status_code == 422
        assert "没有数据行" in headers_only.json()["message"]
