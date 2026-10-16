"""Native governance HTTP acceptance with signed tokens, live identities and isolated data."""

import base64
import io
import json
import os
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal

import httpx
import pytest
from openpyxl import load_workbook
from redis.asyncio import Redis
from sqlalchemy import select

from app.api.v1.health import _readiness_details
from app.core.audit import make_audit_log
from app.core.cache import redis_manager
from app.core.config import settings
from app.core.constants import XLSX_MEDIA_TYPE
from app.core.encryption import filesystem_path
from app.core.security import create_access_token
from app.core.tenant import bypass_tenant_filter
from app.db.session import get_db
from app.main import app
from app.models.audit import AuditLog
from app.models.governance import ConfigRevision, DurableJob, ModelUsage
from app.models.knowledge import KnowledgeDoc
from app.models.meeting import MeetingRecord
from app.models.member import MemberCultivation, MemberProfile
from app.models.qa_session import QASession
from app.models.study import StudyPlan, StudyPlanItem
from app.models.tenant import Tenant
from app.models.user import User, UserRole
from app.services import job_service
from app.services.audit_export import CSV_NAME, GUIDE_NAME, MANIFEST_NAME
from app.services.member_service import _load_scoring_rules
from app.services.qa_record_io import HEADERS, RECORD_SHEET
from tests.integration import test_module6_governance as fixtures

governance_url = fixtures.governance_url
native = fixtures.native
seed = fixtures.seed

pytestmark = pytest.mark.integration


@pytest.fixture
async def http_case(native, monkeypatch, tmp_path):
    _, maker = native
    async with maker() as db:
        school, root, branch, author, admin = await seed(db)
        other_school, _, _, _, other_admin = await seed(db)
    client_redis = Redis.from_url(settings.REDIS_URL, decode_responses=True)
    await client_redis.ping()
    monkeypatch.setattr(redis_manager, "_redis", client_redis)
    monkeypatch.setattr(settings, "AUDIT_SPOOL_DIR", str(tmp_path / "audit"))

    async def database():
        async with maker() as db:
            yield db

    assert get_db not in app.dependency_overrides
    app.dependency_overrides[get_db] = database
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://isolated-test"
        ) as client:
            yield {
                "http": client,
                "maker": maker,
                "school": school,
                "root": root,
                "branch": branch,
                "author": author,
                "admin": admin,
                "other_school": other_school,
                "other_admin": other_admin,
            }
    finally:
        app.dependency_overrides.pop(get_db, None)
        await client_redis.aclose()


def authorization(user, role=None):
    token = create_access_token(str(user.id), str(user.tenant_id), role or user.role)
    return {"Authorization": "Bearer " + token}


async def call(case, actor, method, path, body=None, status=200, headers=None, params=None):
    result = await case["http"].request(
        method,
        "/api/v1" + path,
        json=body,
        params=params,
        headers=headers or authorization(actor),
    )
    assert result.status_code == status, result.text
    if status == 200:
        assert result.json()["code"] == 0, result.text
        return result.json()["data"]
    return result


async def seed_qa_records(case):
    """造问答记录：本租户 3 条（1 条无依据拒答、1 条提问人已删除）+ 另一租户 1 条。"""
    common = {
        "citations": [],
        "warnings": [],
        "retrieved_count": 0,
        "used_count": 0,
        "session_id": uuid.uuid4().hex,
    }
    async with case["maker"]() as db:
        with bypass_tenant_filter():
            db.add_all(
                [
                    QASession(
                        tenant_id=case["school"].id,
                        user_id=case["author"].id,
                        question="转正需要什么材料？",
                        answer="需要提交转正申请等材料。",
                        citations=[{"index": 1, "title": "合成转正规定"}],
                        data_level="internal",
                        has_sufficient_evidence=True,
                        retrieved_count=3,
                        used_count=2,
                        warnings=[],
                        refused=False,
                        session_id=uuid.uuid4().hex,
                        # 版本留痕（SEC-11）：回答产生时记下的版本指纹，随记录一起查得到
                        versions={
                            "prompt": {"qa": "v1", "rewrite": "v1"},
                            "workflow": {"WF-03": "wf-03/1", "WF-04": "wf-04/1"},
                        },
                        created_at=datetime(2026, 3, 1, 1, 0),
                    ),
                    QASession(
                        tenant_id=case["school"].id,
                        user_id=case["author"].id,
                        question="合成拒答问题",
                        answer="暂无依据，无法作答。",
                        data_level="public",
                        has_sufficient_evidence=False,
                        refused=True,
                        created_at=datetime(2026, 3, 2, 1, 0),
                        **common,
                    ),
                    QASession(
                        tenant_id=case["school"].id,
                        user_id=case["author"].id,
                        question="合成已删除问题",
                        answer="提问人已在自己的记录里删除。",
                        data_level="internal",
                        has_sufficient_evidence=True,
                        is_deleted=True,
                        created_at=datetime(2026, 3, 3, 1, 0),
                        **common,
                    ),
                    QASession(
                        tenant_id=case["other_school"].id,
                        user_id=case["other_admin"].id,
                        question="外校合成问题",
                        answer="另一租户的记录，不得跨租户看到。",
                        data_level="internal",
                        has_sufficient_evidence=True,
                        created_at=datetime(2026, 3, 2, 1, 0),
                        **common,
                    ),
                ]
            )
            await db.commit()


async def test_school_admin_http_configuration_identity_conflicts_and_safe_responses(http_case):
    c, admin = http_case, http_case["admin"]
    policy = await call(c, admin, "GET", "/admin/governance")
    assert policy["admission"]["status"] == "pending"
    policy["values"]["prompt_versions"] = {"qa": "v2"}
    update = {
        "expected_version": policy["version"],
        "reason": "isolated test revision",
        "values": policy["values"],
    }
    changed = await call(c, admin, "PATCH", "/admin/governance", update)
    assert changed["version"] == policy["version"] + 1
    await call(c, admin, "PATCH", "/admin/governance", update, status=409)
    history = await call(c, admin, "GET", "/admin/config-history")
    assert history["items"][0]["version"] == changed["version"]
    assert history["items"][0]["actor_id"] == admin.id
    tenants = await call(c, admin, "GET", "/admin/tenants")
    assert {row["id"] for row in tenants["items"]} == {c["school"].id}
    await call(
        c,
        admin,
        "POST",
        "/admin/tenants",
        {"name": "cannot create school", "parent_id": c["school"].id},
        status=403,
    )
    department = await call(
        c,
        admin,
        "POST",
        "/admin/organizations",
        {"name": "isolated department", "org_type": "department", "parent_id": c["root"].id},
    )
    moved = await call(
        c, admin, "PATCH", "/admin/organizations/" + c["branch"].id, {"parent_id": department["id"]}
    )
    assert moved["path"] == department["path"] + "/" + c["branch"].id
    await call(
        c,
        admin,
        "PATCH",
        "/admin/organizations/" + department["id"],
        {"parent_id": c["branch"].id},
        status=422,
    )
    # 开户只允许平台系统管理员（见 ADMINISTRATION.md「系统→创建账号」）。学校管理员在
    # 治理页点"新增用户"会被服务端拒绝，用例把这条 HTTP 边界固定下来。
    await call(
        c,
        admin,
        "POST",
        "/admin/users",
        {
            "username": "native_" + uuid.uuid4().hex,
            "name": "synthetic admin",
            "password": "Isolated_Test_2026!",
            "role": "department_admin",
            "org_unit_id": department["id"],
        },
        status=403,
    )
    # 既有账号的角色/组织一致性校验与读取仍属学校管理员（IDENTITY_MANAGE）职责。
    # 账号由平台管理员开户，这里直接建库准备被测数据。
    user = User(
        id=str(uuid.uuid4()),
        tenant_id=c["school"].id,
        username="native_" + uuid.uuid4().hex,
        name="synthetic admin",
        password_hash="synthetic",
        role=UserRole.DEPARTMENT_ADMIN,
        org_unit_id=department["id"],
        is_active=True,
    )
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            db.add(user)
            await db.commit()
    await call(c, admin, "PATCH", "/admin/users/" + user.id, {"role": "school_admin"}, status=422)
    users = await call(c, admin, "GET", "/admin/users")
    assert next(row for row in users["items"] if row["id"] == user.id)["role"] == "department_admin"
    await call(
        c, admin, "GET", "/admin/organizations?tenant_id=" + c["other_school"].id, status=403
    )
    await call(
        c,
        admin,
        "GET",
        "/admin/metrics",
        status=403,
        headers=authorization(admin, UserRole.SYSTEM_ADMIN),
    )


async def test_same_signed_token_observes_role_user_and_ancestor_revocation(http_case):
    c, admin = http_case, http_case["admin"]
    # 开户限平台系统管理员，这里直接建库准备待撤销的学校管理员账号。
    actor = User(
        id=str(uuid.uuid4()),
        tenant_id=c["school"].id,
        username="native_" + uuid.uuid4().hex,
        name="synthetic school admin",
        password_hash="synthetic",
        role=UserRole.SCHOOL_ADMIN,
        org_unit_id=c["root"].id,
        is_active=True,
    )
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            db.add(actor)
            await db.commit()
    same_token = authorization(actor)
    await call(c, actor, "GET", "/admin/users", headers=same_token)
    await call(
        c,
        admin,
        "PATCH",
        "/admin/users/" + actor.id,
        {"role": "member", "org_unit_id": c["branch"].id},
    )
    await call(c, actor, "GET", "/admin/users", headers=same_token, status=403)
    await call(c, actor, "GET", "/admin/jobs", headers=same_token)
    await call(c, admin, "PATCH", "/admin/users/" + actor.id, {"is_active": False})
    await call(c, actor, "GET", "/admin/jobs", headers=same_token, status=403)
    await call(c, admin, "PATCH", "/admin/users/" + admin.id, {"is_active": False}, status=409)
    author_token = authorization(c["author"])
    await call(c, c["author"], "GET", "/admin/jobs", headers=author_token)
    await call(c, admin, "PATCH", "/admin/organizations/" + c["branch"].id, {"is_active": False})
    await call(c, c["author"], "GET", "/admin/jobs", headers=author_token, status=403)


async def test_http_reviews_reject_self_stale_cross_school_and_freeze_current_evidence(http_case):
    c = http_case
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            profile = MemberProfile(
                tenant_id=c["school"].id,
                name="synthetic candidate",
                org_name=c["branch"].name,
                org_unit_id=c["branch"].id,
            )
            db.add(profile)
            await db.flush()
            for category in _load_scoring_rules()["weights"]:
                db.add(
                    MemberCultivation(
                        tenant_id=c["school"].id,
                        profile_id=profile.id,
                        org_unit_id=c["branch"].id,
                        category=category,
                        score=Decimal("50"),
                        period="2026",
                        source_note="synthetic test evidence",
                    )
                )
            await db.commit()
    body = {"source_kind": "score", "source_id": profile.id, "year": 2026}
    review = await call(c, c["author"], "POST", "/admin/reviews", body)
    decision = {
        "decision": "approved",
        "source_version": review["source_version"],
        "opinion": "synthetic human check",
    }
    path = "/admin/reviews/" + review["id"]
    await call(c, c["author"], "POST", path + "/decision", decision, status=403)
    await call(
        c,
        c["admin"],
        "POST",
        path + "/decision",
        {**decision, "source_version": "f" * 64},
        status=409,
    )
    await call(c, c["other_admin"], "GET", path, status=404)
    approved = await call(c, c["admin"], "POST", path + "/decision", decision)
    assert approved["status"] == "approved"
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            evidence = (
                (
                    await db.execute(
                        select(MemberCultivation).where(MemberCultivation.profile_id == profile.id)
                    )
                )
                .scalars()
                .first()
            )
            evidence.score = Decimal("80")
            await db.commit()
    stale = await call(c, c["admin"], "GET", path)
    assert stale["stale"] and stale["snapshot"] is None
    fresh = await call(c, c["author"], "POST", "/admin/reviews", body)
    assert fresh["id"] != review["id"]
    listing = await call(c, c["admin"], "GET", "/admin/reviews?status=pending")
    assert listing["total"] == 1
    missing = "/admin/reviews/" + str(uuid.uuid4())
    await call(c, c["admin"], "GET", missing, status=404)


async def test_jobs_are_actor_private_and_operations_use_only_school_durable_evidence(http_case):
    c = http_case
    body = {"kind": "member_reminders", "dedup_key": "bounded-http-retry"}
    job = await call(c, c["author"], "POST", "/admin/jobs", body)
    assert (await call(c, c["author"], "POST", "/admin/jobs", body))["id"] == job["id"]
    own = await call(c, c["author"], "GET", "/admin/jobs")
    assert [row["id"] for row in own["items"]] == [job["id"]]
    assert (await call(c, c["admin"], "GET", "/admin/jobs"))["items"] == []
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            db.add(
                ModelUsage(
                    tenant_id=c["school"].id,
                    request_id=uuid.uuid4().hex,
                    config_version=0,
                    model_type="llm",
                    model_name="synthetic-only",
                    result="success",
                    duration_seconds=0.5,
                )
            )
            for _ in range(5):
                db.add(
                    AuditLog(
                        tenant_id=c["school"].id,
                        user_id=c["admin"].id,
                        user_name="synthetic",
                        action="export",
                        resource_type="synthetic",
                        request_id=uuid.uuid4().hex,
                        result="denied",
                        data_level="internal",
                    )
                )
            await db.commit()
    ops = await call(c, c["admin"], "GET", "/admin/operations")
    assert ops["jobs"]["queued"] == 1
    assert ops["admission"]["status"] == "pending"
    assert any(row["calls"] == 1 and row["type"] == "llm" for row in ops["model_usage"])
    assert any(row["result"] == "denied" and row["count"] >= 5 for row in ops["alerts"])
    foreign = await call(c, c["other_admin"], "GET", "/admin/operations")
    assert foreign["jobs"] == {} and foreign["model_usage"] == []
    assert await job_service.run_once()
    finished = await call(c, c["author"], "GET", "/admin/jobs")
    assert finished["items"][0]["status"] == "done"
    assert finished["items"][0]["result"] == {"processed": 0}
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            completed = (
                (
                    await db.execute(
                        select(AuditLog).where(
                            AuditLog.resource_id == job["id"], AuditLog.action == "job_completed"
                        )
                    )
                )
                .scalars()
                .all()
            )
    assert len(completed) == 1


async def test_run_logs_show_attempts_failures_and_scan_results(http_case):
    """T-SYS-06E：运行日志要回答"这个租户最近跑了什么、失败在哪一步"（admin.ops，租户内可见）。"""
    c = http_case
    # 支部书记没有 admin.ops：运行日志不对其开放（本人任务状态在「任务台账」另有口径）
    await call(c, c["author"], "GET", "/admin/run-logs", status=403)
    job_body = {"kind": "member_reminders", "dedup_key": "run-log-ok"}
    job = await call(c, c["author"], "POST", "/admin/jobs", job_body)
    # 一个必然失败的任务：文档向量化要 knowledge.manage，支部书记没有 → 403 直接判定失败
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            failing = DurableJob(
                tenant_id=c["school"].id,
                actor_id=c["author"].id,
                kind="embedding",
                dedup_key="run-log-fail",
                payload={"document_ids": [uuid.uuid4().hex]},
                status="queued",
                attempts=0,
                max_attempts=3,
                available_at=datetime(2020, 1, 1),
            )
            db.add(failing)
            await db.commit()
            failing_id = str(failing.id)
    assert await job_service.run_once()
    assert await job_service.run_once()
    page = await call(c, c["admin"], "GET", "/admin/run-logs")
    assert page["scope_note"] == "租户 " + c["school"].id
    ids = {row["resource_id"] for row in page["items"]}
    assert {job["id"], failing_id} <= ids
    completed = next(row for row in page["items"] if row["event"] == "job_completed")
    assert completed["resource_id"] == job["id"]
    assert completed["event_text"] == "执行完成" and completed["result"] == "success"
    assert completed["kind"] == "member_reminders" and completed["processed"] == 0
    # 耗时是排障的关键：能区分"秒挂"和"跑到超时"
    assert isinstance(completed["duration_ms"], int) and completed["duration_ms"] >= 0
    failed = next(row for row in page["items"] if row["event"] == "job_attempt_failed")
    assert failed["resource_id"] == failing_id
    assert failed["event_text"] == "执行失败" and failed["result"] == "denied"
    assert failed["kind"] == "embedding" and failed["status"] == "failed"
    assert failed["attempt"] == 1 and failed["error_code"] == "HTTPException"
    assert isinstance(failed["duration_ms"], int)
    # 入队也在日志里；任务类型以任务表为准（入队审计本身是摘要口径，拿不到明文类型）
    enqueued = next(row for row in page["items"] if row["event"] == "job_enqueue")
    assert enqueued["resource_id"] == job["id"] and enqueued["kind"] == "member_reminders"
    # 事件 / 结果 / 时间窗筛选；未知取值给可执行错误而不是空结果
    only_failed = await call(
        c, c["admin"], "GET", "/admin/run-logs", params={"event": "job_attempt_failed"}
    )
    assert [row["resource_id"] for row in only_failed["items"]] == [failing_id]
    assert only_failed["scope_note"].endswith("事件 执行失败")
    # 任务类型按任务表口径筛：只出这类任务的事件（定时扫描不属于某一条任务，故不出现）
    by_kind = await call(c, c["admin"], "GET", "/admin/run-logs", params={"kind": "embedding"})
    assert [row["resource_id"] for row in by_kind["items"]] == [failing_id]
    assert "任务类型 embedding" in by_kind["scope_note"]
    assert (
        await call(
            c,
            c["admin"],
            "GET",
            "/admin/run-logs",
            params={"event": "job_attempt_failed", "result": "success"},
        )
    )["total"] == 0
    for bad in ({"event": "不存在的事件"}, {"kind": "unknown"}):
        await call(c, c["admin"], "GET", "/admin/run-logs", params=bad, status=422)
    wide = await call(
        c, c["admin"], "GET", "/admin/run-logs", params={"start_at": "2000-01-01T00:00:00+08:00"}
    )
    assert wide["total"] == page["total"]
    assert (
        await call(
            c, c["admin"], "GET", "/admin/run-logs", params={"end_at": "2000-01-01T00:00:00+08:00"}
        )
    )["total"] == 0
    await call(
        c,
        c["admin"],
        "GET",
        "/admin/run-logs",
        params={"start_at": "2027-01-01T00:00:00+08:00", "end_at": "2020-01-01T00:00:00+08:00"},
        status=422,
    )
    # 跨租户不可见：外校管理员看不到本校的运行日志
    assert (await call(c, c["other_admin"], "GET", "/admin/run-logs"))["total"] == 0


async def test_version_trail_shows_prompt_workflow_and_knowledge_versions(http_case):
    """T-SEC-11：提示词 / 工作流 / 知识库三个版本都要可查、可追溯（admin.audit，租户内）。"""
    c = http_case
    # 支部书记没有 admin.audit：版本留痕与问答记录、审计查询同一枚权限
    await call(c, c["author"], "GET", "/admin/version-trail", status=403)
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            db.add_all(
                [
                    # 三次配置提交，只有两次动了提示词：第 2 次只调了别的字段
                    ConfigRevision(
                        tenant_id=c["school"].id,
                        version=1,
                        actor_id=c["admin"].id,
                        reason="初次登记",
                        values={"prompt_versions": {"qa": "v1"}},
                    ),
                    ConfigRevision(
                        tenant_id=c["school"].id,
                        version=2,
                        actor_id=c["admin"].id,
                        reason="只调了存储期限",
                        values={"prompt_versions": {"qa": "v1"}, "retention": {"audit_days": 400}},
                    ),
                    ConfigRevision(
                        tenant_id=c["school"].id,
                        version=3,
                        actor_id=c["admin"].id,
                        reason="问答提示词升版",
                        values={"prompt_versions": {"qa": "v2"}},
                    ),
                ]
            )
            document = KnowledgeDoc(
                tenant_id=c["school"].id,
                doc_id="vt-" + uuid.uuid4().hex,
                file_name="synthetic.txt",
                title="版本留痕合成制度",
                issuer="合成单位",
                level="school",
                visibility="school",
                security_level="public",
                effective_date=date(2020, 1, 1),
                status="effective",
                tags=[],
                doc_metadata={"content_revision": 3, "org_unit_id": c["root"].id},
            )
            db.add(document)
            await db.flush()
            doc_pk, doc_business_id = str(document.id), document.doc_id

            def audit(action, *, request_id, **values):
                row = make_audit_log(
                    tenant_id=c["school"].id,
                    user_id=c["admin"].id,
                    user_name=str(c["admin"].id),
                    action=action,
                    resource_type="knowledge_doc",
                    resource_id=doc_pk,
                    request_id=request_id,
                    result="success",
                    **values,
                )
                return row

            created = audit("create", request_id="vt-create", new_value={"chunk_count": 4})
            created.created_at = datetime(2026, 3, 1, 1, 0)
            updated = audit(
                "update_content",
                request_id="vt-update",
                old_value={"revision": 2},
                new_value={"revision": 3, "chunk_count": 5},
            )
            updated.created_at = datetime(2026, 3, 1, 2, 0)
            # 只改元数据：不留修订号，不算版本变更，不该出现在留痕里
            metadata_only = audit(
                "update_metadata", request_id="vt-meta", new_value={"title": "改了标题"}
            )
            metadata_only.created_at = datetime(2026, 3, 1, 3, 0)
            db.add_all([created, updated, metadata_only])
            await db.commit()

    data = await call(c, c["admin"], "GET", "/admin/version-trail")
    # 提示词：当前生效值来自租户配置（未配置回落出厂默认），默认值也明确给出
    assert data["prompt"]["current"] == {"qa": "v1", "rewrite": "v1"}
    assert data["prompt"]["defaults"] == {"qa": "v1", "rewrite": "v1"}
    history = data["prompt"]["history"]
    assert [entry["config_version"] for entry in history] == [3, 1]
    assert history[0]["reason"] == "问答提示词升版" and history[0]["actor"] == "测试审核人"
    assert [
        (change["kind"], change["before"], change["after"]) for change in history[0]["changes"]
    ] == [("qa", "v1", "v2")]
    # 首次登记从"无"变成当时的取值：这一条本身也是留痕，不能只剩"最近一次改动"
    assert history[1]["changes"][0]["before"] is None
    assert c["school"].id in data["prompt"]["scope_note"]
    # 工作流：契约版本是代码常量，只给现行值，不编造数据库里没有的变更历史
    assert {item["workflow"]: item["version"] for item in data["workflow"]["items"]} == {
        "WF-03": "wf-03/1",
        "WF-04": "wf-04/1",
    }
    # 知识库：入库即第 1 版；内容更新给出前后版本；只改元数据不出现
    knowledge = data["knowledge"]
    assert [row["event"] for row in knowledge["items"]] == ["update_content", "create"]
    assert knowledge["total"] == 2
    assert knowledge["items"][0]["previous_revision"] == 2 and knowledge["items"][0]["revision"] == 3
    assert knowledge["items"][0]["event_text"] == "内容更新"
    assert knowledge["items"][0]["doc_id"] == doc_business_id
    assert knowledge["items"][0]["title"] == "版本留痕合成制度"
    assert knowledge["items"][0]["actor"] == "测试审核人"
    assert knowledge["items"][1]["revision"] == 1 and knowledge["items"][1]["previous_revision"] is None
    assert c["school"].id in knowledge["scope_note"]
    # 跨租户不可见：外校管理员既看不到本校的提示词历史，也看不到知识库修订
    foreign = await call(c, c["other_admin"], "GET", "/admin/version-trail")
    assert foreign["prompt"]["history"] == []
    assert foreign["knowledge"]["total"] == 0


async def test_production_readiness_keeps_unconfirmed_school_and_models_unready(
    http_case, native, monkeypatch
):
    import app.db.session as sessions

    engine, maker = native
    monkeypatch.setattr(sessions, "engine", engine)
    monkeypatch.setattr(sessions, "async_session_maker", maker)
    monkeypatch.setattr(settings, "ENV", "production")
    monkeypatch.setattr(settings, "SECRET_KEY", "synthetic-readiness-only-secret-2026")
    monkeypatch.setenv(
        "FILE_ENCRYPTION_KEYS", '{"acceptance":"' + base64.b64encode(os.urandom(32)).decode() + '"}'
    )
    monkeypatch.setenv("FILE_ENCRYPTION_ACTIVE_KEY", "acceptance")
    # 本用例断言的是「模型尚未确认时不得报 healthy」，而 model_readiness() 读的是进程级
    # 全局注册表；本地 .env 若配置了真实密钥，只要同进程内跑过一次应用 lifespan，注册表
    # 就会带上云端模型。这里显式固定注册状态，避免用例之间通过该单例互相影响。
    import app.llm.init_models as init_models

    monkeypatch.setattr(
        init_models,
        "model_readiness",
        lambda: {"llm": "missing", "embedding": "missing", "reranker": "missing"},
    )
    response = await http_case["http"].get("/api/v1/health/ready")
    assert response.status_code == 503, response.text
    # 公开探针只暴露整体状态；依赖明细只对平台管理员开放（/admin/health/ready）。
    # 这里同时断言"不泄露明细"，再用同一实现核对未就绪的具体原因。
    assert response.json() == {"status": "not_ready"}
    details = json.loads((await _readiness_details()).body)
    checks = details["checks"]
    for dependency in ("database", "redis", "schema", "encryption", "audit_storage", "jwt_secret"):
        assert checks[dependency] == "healthy"
    assert checks["school_admission"] == "configuration_required"
    assert checks["runtime_role"] == "configuration_required"
    assert checks["models"] != "healthy"


async def test_production_readiness_ignores_recent_model_traffic(http_case, native, monkeypatch):
    """模型已配置但近期无人提问时，就绪不得因此翻转（连通性只作观察项）。"""
    import app.db.session as sessions
    import app.llm.init_models as init_models

    engine, maker = native
    monkeypatch.setattr(sessions, "engine", engine)
    monkeypatch.setattr(sessions, "async_session_maker", maker)
    monkeypatch.setattr(settings, "ENV", "production")
    monkeypatch.setattr(settings, "SECRET_KEY", "synthetic-readiness-cold-model-secret-2026")
    monkeypatch.setenv(
        "FILE_ENCRYPTION_KEYS", '{"acceptance":"' + base64.b64encode(os.urandom(32)).decode() + '"}'
    )
    monkeypatch.setenv("FILE_ENCRYPTION_ACTIVE_KEY", "acceptance")
    monkeypatch.setattr(
        init_models,
        "model_readiness",
        lambda: {"llm": "configured", "embedding": "configured", "reranker": "configured"},
    )
    response = await http_case["http"].get("/api/v1/health/ready")
    details = json.loads((await _readiness_details()).body)
    checks = details["checks"]
    # 配置齐全即 model 健康；「最近 15 分钟有成功调用」只体现在 model_warm 观察项上。
    assert checks["models"] == "healthy"
    assert checks["model_warm"] == "unverified"
    # 学校效力确认与运行角色仍未确认，整体依旧未就绪——但原因不应是「模型冷」。
    assert response.status_code == 503, response.text
    assert checks["school_admission"] == "configuration_required"


async def test_encrypted_original_versions_survive_update_and_reject_foreign_downloads(
    http_case, monkeypatch, tmp_path
):
    c = http_case
    monkeypatch.setenv("FILE_STORAGE_DIR", str(tmp_path / "originals"))
    identifier = "M6-HTTP-" + uuid.uuid4().hex
    path = "/api/v1/knowledge-docs/" + identifier
    initial = "第一条 合成公开验收规定。\n第二条 原件应保留版本。".encode()
    revised = "第一条 合成公开验收新规定。\n第二条 原件应保留历史版本。".encode()
    form = {
        "doc_id": identifier,
        "file_name": "initial.txt",
        "title": "合成公开验收规定",
        "issuer": "合成机构",
        "level": "school",
        "visibility": "school",
        "effective_date": "2026-01-01",
        "tags": "验收",
        "security_level": "public",
    }
    headers = authorization(c["admin"])
    created = await c["http"].post(
        "/api/v1/knowledge-docs",
        data=form,
        files={"file": ("initial.txt", initial)},
        headers=headers,
    )
    assert created.status_code == 200, created.text
    updated = await c["http"].put(
        path + "/content",
        data={"expected_revision": "1"},
        files={"file": ("revised.txt", revised)},
        headers=headers,
    )
    assert updated.status_code == 200, updated.text
    for revision, expected in ((1, initial), (2, revised)):
        download = await c["http"].get(path + f"/file?revision={revision}", headers=headers)
        assert download.status_code == 200 and download.content == expected
        foreign = await c["http"].get(
            path + f"/file?revision={revision}", headers=authorization(c["other_admin"])
        )
        assert foreign.status_code == 404
    conflict = await c["http"].put(
        path + "/content",
        data={"expected_revision": "1"},
        files={"file": ("stale.txt", initial)},
        headers=headers,
    )
    assert conflict.status_code == 409
    for filename, content, expected in (("empty.txt", b"", 422), ("bad.exe", b"binary", 415)):
        invalid = await c["http"].put(
            path + "/content",
            data={"expected_revision": "2"},
            files={"file": (filename, content)},
            headers=headers,
        )
        assert invalid.status_code == expected
    missing = await c["http"].get(path + "/file?revision=3", headers=headers)
    assert missing.status_code == 404
    encrypted_files = list((tmp_path / "originals").rglob("*.bin"))
    assert len(encrypted_files) == 2
    # 存储根 + 内容寻址文件名可能超过 Windows MAX_PATH，读取沿用应用的路径口径。
    for item in encrypted_files:
        stored = filesystem_path(item).read_bytes()
        assert initial not in stored and revised not in stored


async def test_platform_scope_school_registration_and_process_metrics_require_real_platform(
    http_case,
):
    c = http_case
    platform_id = str(uuid.uuid4())
    platform_actor = User(
        id=str(uuid.uuid4()),
        tenant_id=platform_id,
        username="platform_" + uuid.uuid4().hex,
        name="synthetic platform admin",
        password_hash="synthetic",
        role=UserRole.SYSTEM_ADMIN,
        is_active=True,
    )
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            db.add(
                Tenant(
                    id=platform_id,
                    tenant_id=platform_id,
                    tenant_type="platform",
                    name="synthetic platform",
                    path="/" + platform_id,
                    config={},
                    is_active=True,
                )
            )
            db.add(platform_actor)
            await db.commit()
    school = await call(
        c,
        platform_actor,
        "POST",
        "/admin/tenants",
        {"name": "synthetic school", "parent_id": platform_id},
    )
    department = await call(
        c,
        platform_actor,
        "POST",
        "/admin/tenants",
        {"name": "synthetic department", "tenant_type": "department", "parent_id": school["id"]},
    )
    await call(
        c,
        platform_actor,
        "POST",
        "/admin/tenants",
        {"name": "invalid branch", "tenant_type": "branch", "parent_id": platform_id},
        status=422,
    )
    listing = await call(c, platform_actor, "GET", "/admin/tenants")
    assert {row["id"] for row in listing["items"]} == {platform_id, school["id"], department["id"]}
    await call(c, platform_actor, "PATCH", "/admin/tenants/" + school["id"], {"is_active": False})
    await call(
        c,
        platform_actor,
        "PATCH",
        "/admin/tenants/" + platform_id,
        {"is_active": False},
        status=409,
    )
    metrics = await c["http"].get("/api/v1/admin/metrics", headers=authorization(platform_actor))
    assert metrics.status_code == 200 and "text/plain" in metrics.headers["content-type"]
    await call(c, c["admin"], "GET", "/admin/metrics", status=403)


async def test_governed_statistics_do_not_count_drafts_or_other_schools(http_case):
    c = http_case
    year = 2026
    participants = [
        {"participant_id": "synthetic-a", "name": "合成人员甲", "attended": True},
        {"participant_id": "synthetic-b", "name": "合成人员乙", "attended": False},
    ]
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            records = []
            for school, org, kind, status in (
                (c["school"], c["root"], "center_group", "approved"),
                (c["school"], c["root"], "center_group", "draft"),
                (c["school"], c["branch"], "party_lecture", "approved"),
                (c["school"], c["branch"], "party_lecture", "pending"),
                (c["other_school"], c["root"], "center_group", "approved"),
            ):
                record = MeetingRecord(
                    tenant_id=school.id,
                    org_unit_id=org.id,
                    activity_type=kind,
                    title="合成统计证据",
                    scheduled_on=date(year, 1, 2),
                    held_on=date(year, 1, 2),
                    review_status=status,
                    reviewed_by=c["admin"].id if status == "approved" else None,
                    host="合成主持人",
                    participants=participants,
                    minutes={
                        "learning_points": "学习",
                        "consensus": "共识",
                        "requirements": "要求",
                    },
                )
                db.add(record)
                records.append(record)
            plan = StudyPlan(
                tenant_id=c["school"].id,
                org_unit_id=c["root"].id,
                year=year,
                title="合成年度计划",
                responsible="合成负责人",
                review_status="approved",
                reviewed_by=c["admin"].id,
            )
            db.add(plan)
            await db.flush()
            for index, record in enumerate(records[:2]):
                db.add(
                    StudyPlanItem(
                        tenant_id=c["school"].id,
                        plan_id=plan.id,
                        topic="合成学习",
                        scheduled_on=date(year, 1, 2),
                        responsible="合成负责人",
                        meeting_record_id=record.id,
                    )
                )
            await db.commit()
    metrics = await call(
        c, c["admin"], "GET", "/study/metrics?year=2026&org_unit_id=" + c["root"].id
    )
    assert metrics["learning_count"]["value"] == 1
    assert metrics["plan_completion_rate"]["value"] == 0.5
    assert metrics["attendance_rate"]["value"] == 0.5
    assert len(metrics["evidence"]) == 1 and metrics["missing"]
    assert metrics["evidence"][0]["activity_id"] == records[0].id
    meetings = await call(
        c, c["admin"], "GET", "/meeting/stats?year=2026&org_unit_id=" + c["branch"].id
    )
    assert meetings["attendance_rate"]["value"] == 0.5
    assert len(meetings["evidence"]) == 1 and meetings["missing"]
    await call(
        c,
        c["admin"],
        "GET",
        "/study/metrics?year=2026&org_unit_id=" + str(uuid.uuid4()),
        status=403,
    )
    empty = await call(c, c["other_admin"], "GET", "/study/metrics?year=2026")
    assert empty["learning_count"]["value"] is None


async def test_qa_records_central_view_is_permission_scoped_and_filterable(http_case):
    """T-SYS-06：问答记录集中查看受 admin.audit 管辖、租户内可见、条件与时间都可核对。"""
    c = http_case
    await seed_qa_records(c)
    # 支部书记没有 admin.audit：集中查看不对其开放（本人档案另有 /qa/sessions 口径）
    await call(c, c["author"], "GET", "/qa/records", status=403)
    page = await call(c, c["admin"], "GET", "/qa/records")
    assert page["total"] == 2
    assert page["scope_note"] == "不含本人已删除的记录"
    assert {row["question"] for row in page["items"]} == {"转正需要什么材料？", "合成拒答问题"}
    # 每条都要能自证：谁问的、什么级别、有没有依据、引用了几处
    answered = next(row for row in page["items"] if row["question"] == "转正需要什么材料？")
    assert answered["asker"] == "测试提交人" and answered["username"] == c["author"].username
    assert answered["data_level"] == "internal" and answered["has_sufficient_evidence"] is True
    assert answered["citation_count"] == 1 and answered["refused"] is False
    assert answered["is_deleted"] is False
    # 版本留痕（SEC-11）：这条答案是哪个版本产出的，要跟记录一起查得到
    assert answered["versions"]["prompt"] == {"qa": "v1", "rewrite": "v1"}
    assert answered["versions"]["workflow"]["WF-03"] == "wf-03/1"
    # 上线前的历史记录没有版本指纹：给空对象而不是按当前配置倒推
    refused = next(row for row in page["items"] if row["question"] == "合成拒答问题")
    assert refused["versions"] == {}
    assert (
        await call(c, c["admin"], "GET", "/qa/records?refused=true")
    )["items"][0]["question"] == "合成拒答问题"
    assert (
        await call(c, c["admin"], "GET", "/qa/records?has_evidence=false")
    )["total"] == 1
    assert (await call(c, c["admin"], "GET", "/qa/records?keyword=转正"))["total"] == 1
    assert (await call(c, c["admin"], "GET", "/qa/records?data_level=sensitive"))["items"] == []
    assert (await call(c, c["admin"], "GET", "/qa/records?user_id=" + str(uuid.uuid4())))["total"] == 0
    # 软删除只是不再显示；显式要求时仍能看到，否则删除就成了审计盲区
    included = await call(c, c["admin"], "GET", "/qa/records?include_deleted=true")
    assert included["total"] == 3
    assert "含本人已删除的记录" in included["scope_note"]
    # 时间按北京时间传入（+08:00）：09:00 正好命中那条 01:00 UTC 的记录，
    # 若被当成本地/UTC 朴素时间就会漏掉它。
    window = await call(
        c,
        c["admin"],
        "GET",
        "/qa/records",
        params={"start_at": "2026-03-02T09:00:00+08:00"},
    )
    assert window["total"] == 1 and window["items"][0]["question"] == "合成拒答问题"
    await call(
        c,
        c["admin"],
        "GET",
        "/qa/records",
        params={"start_at": "2026-03-02T00:00:00+08:00", "end_at": "2026-03-01T00:00:00+08:00"},
        status=422,
    )
    # 跨租户：本平台管理员之外的账号不得借 tenant_id 越界
    await call(c, c["admin"], "GET", "/qa/records?tenant_id=" + c["other_school"].id, status=403)
    foreign = await call(c, c["other_admin"], "GET", "/qa/records")
    assert [row["question"] for row in foreign["items"]] == ["外校合成问题"]


async def test_qa_record_export_carries_its_scope_and_writes_audit(http_case, monkeypatch):
    """T-SYS-06：导出的 Excel 必须自证筛选口径，且导出动作留审计、超限给可执行提示。"""
    c = http_case
    await seed_qa_records(c)
    denied = await c["http"].get(
        "/api/v1/qa/records/export", headers=authorization(c["author"])
    )
    assert denied.status_code == 403
    response = await c["http"].get(
        "/api/v1/qa/records/export", params={"keyword": "转正"}, headers=authorization(c["admin"])
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert "qa-records.xlsx" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook[RECORD_SHEET]
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.max_row == 2
    # 版本留痕（SEC-11）：台账里"这条答案用了哪个版本"要跟着记录走，不能只在界面上有
    assert sheet.cell(row=2, column=8).value == "提示词：问答提示词 v1 / 追问改写 v1；工作流：WF-03 wf-03/1 / WF-04 wf-04/1"
    assert sheet.cell(row=2, column=12).value == "转正需要什么材料？"
    guide = "\n".join(str(cell.value) for cell in workbook["导出说明"]["A"] if cell.value)
    assert "导出人：测试审核人" in guide
    assert "导出条数：1" in guide
    assert "筛选口径：不含本人已删除的记录；问题或回答含「转正」" in guide
    # "未留痕"要有解释，否则复核时会把空白当成漏记
    assert "「版本」为本次回答的版本留痕" in guide
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            logs = (
                (
                    await db.execute(
                        select(AuditLog).where(
                            AuditLog.action == "export", AuditLog.resource_type == "qa_record"
                        )
                    )
                )
                .scalars()
                .all()
            )
    assert len(logs) == 1
    assert logs[0].tenant_id == c["school"].id and logs[0].user_id == c["admin"].id
    # 超过上限：提示分批导出，而不是生成一份打不开的大文件
    monkeypatch.setattr("app.api.v1.qa.MAX_EXPORT_ROWS", 1)
    oversized = await c["http"].get("/api/v1/qa/records/export", headers=authorization(c["admin"]))
    assert oversized.status_code == 413, oversized.text
    assert "分批导出" in oversized.json()["message"]


async def seed_audit_logs(case):
    """造审计记录：本租户 3 条（含 1 条拒绝）+ 另一租户 1 条同动作记录，用于验证跨租户边界。"""
    rows = (
        ("create", "user", "success", None, datetime(2026, 3, 1, 1, 0), case["school"].id),
        ("delete", "member", "denied", "权限或业务校验未通过", datetime(2026, 3, 1, 2, 0), case["school"].id),
        ("update", "member", "success", None, datetime(2026, 3, 1, 3, 0), case["school"].id),
        ("delete", "member", "success", None, datetime(2026, 3, 2, 2, 0), case["other_school"].id),
    )
    async with case["maker"]() as db:
        with bypass_tenant_filter():
            for action, resource_type, result, error, created_at, tenant_id in rows:
                row = make_audit_log(
                    tenant_id=str(tenant_id),
                    user_id=str(case["admin"].id),
                    user_name="合成审计对象",
                    action=action,
                    resource_type=resource_type,
                    resource_id="合成对象",
                    request_id=uuid.uuid4().hex,
                    result=result,
                    data_level="sensitive",
                    ip_address="10.1.2.3",
                    new_value={"note": "合成审计"},
                    error_message=error,
                )
                row.created_at = created_at
                db.add(row)
            await db.commit()


async def test_audit_query_filters_and_cross_tenant_scope(http_case):
    """T-SYS-05E：审计高级检索按动作 / 级别 / 时间可核对，跨租户仅平台管理员。"""
    c = http_case
    await seed_audit_logs(c)
    page = await call(c, c["admin"], "GET", "/admin/audit-logs", params={"action": "delete"})
    assert page["total"] == 1
    assert page["items"][0]["resource_type"] == "member"
    assert page["items"][0]["result"] == "denied"
    assert page["items"][0]["user_name"] == "合成审计对象"
    assert page["scope_note"] == "租户 " + c["school"].id + "；动作 delete"
    # 数据级别筛选与动作叠加：中间件写的请求级审计不会伪装成我们的合成记录
    assert (
        await call(
            c,
            c["admin"],
            "GET",
            "/admin/audit-logs",
            params={"action": "delete", "data_level": "sensitive"},
        )
    )["total"] == 1
    assert (
        await call(
            c,
            c["admin"],
            "GET",
            "/admin/audit-logs",
            params={"action": "delete", "data_level": "public"},
        )
    )["total"] == 0
    # 时间按北京时间传入：02:00 UTC 的记录用 09:00+08:00 起查命中，11:00+08:00 起查落空
    assert (
        await call(
            c,
            c["admin"],
            "GET",
            "/admin/audit-logs",
            params={"action": "delete", "start_at": "2026-03-01T09:00:00+08:00"},
        )
    )["total"] == 1
    assert (
        await call(
            c,
            c["admin"],
            "GET",
            "/admin/audit-logs",
            params={"action": "delete", "start_at": "2026-03-01T11:00:00+08:00"},
        )
    )["total"] == 0
    await call(
        c,
        c["admin"],
        "GET",
        "/admin/audit-logs",
        params={"start_at": "2026-03-02T00:00:00+08:00", "end_at": "2026-03-01T00:00:00+08:00"},
        status=422,
    )
    # 跨租户检索只有真实平台管理员可用；外校管理员只看得到自己租户的记录
    await call(
        c,
        c["admin"],
        "GET",
        "/admin/audit-logs",
        params={"tenant_id": c["other_school"].id},
        status=403,
    )
    foreign = await call(c, c["other_admin"], "GET", "/admin/audit-logs", params={"action": "delete"})
    assert foreign["total"] == 1
    assert foreign["items"][0]["request_id"] != page["items"][0]["request_id"]


async def test_audit_export_package_is_self_documenting_and_capped(http_case, monkeypatch):
    """T-SYS-05E / T-NFR-08：台账导出必须自证筛选口径、留审计、超限给可执行提示。"""
    c = http_case
    await seed_audit_logs(c)
    denied = await c["http"].get(
        "/api/v1/admin/audit-logs/export", headers=authorization(c["author"])
    )
    assert denied.status_code == 403
    response = await c["http"].get(
        "/api/v1/admin/audit-logs/export",
        params={"action": "delete"},
        headers=authorization(c["admin"]),
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/zip")
    assert "audit-logs.zip" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert set(archive.namelist()) == {CSV_NAME, GUIDE_NAME, MANIFEST_NAME}
        manifest = json.loads(archive.read(MANIFEST_NAME).decode("utf-8"))
        csv_text = archive.read(CSV_NAME).decode("utf-8-sig")
        guide = archive.read(GUIDE_NAME).decode("utf-8")
    assert manifest["导出条数"] == 1
    assert manifest["导出人"] == "测试审核人（" + c["admin"].username + "）"
    assert manifest["筛选口径"] == "租户 " + c["school"].id + "；动作 delete"
    assert "2026-03-01 10:00" in csv_text  # 02:00 UTC → 10:00 北京时间
    assert "member" in csv_text and "权限或业务校验未通过" in csv_text
    assert "导出条数：1" in guide
    # 导出动作本身必须留痕：谁在什么时候按什么口径导了什么
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            logs = (
                (
                    await db.execute(
                        select(AuditLog).where(
                            AuditLog.action == "export", AuditLog.resource_type == "audit_log"
                        )
                    )
                )
                .scalars()
                .all()
            )
    assert len(logs) == 1
    assert logs[0].tenant_id == c["school"].id and logs[0].user_id == c["admin"].id
    assert logs[0].data_level == "sensitive"
    # 审计摘要口径：业务内容只留摘要（sha256 + 字节数），可核对但不可反查
    assert logs[0].new_value["bytes"] > 0 and logs[0].new_value["sha256"]
    # 超过上限：提示分批导出，而不是生成一份打不开的大文件
    monkeypatch.setattr("app.api.v1.admin.MAX_EXPORT_ROWS", 1)
    oversized = await c["http"].get(
        "/api/v1/admin/audit-logs/export",
        params={"resource_type": "member"},
        headers=authorization(c["admin"]),
    )
    assert oversized.status_code == 413, oversized.text
    assert "分批导出" in oversized.json()["message"]


async def test_member_scheduler_repeats_without_duplicate_reminders_and_keeps_scope(
    http_case, monkeypatch
):
    from app.models.member import MemberReminder
    from app.services import member_scheduler

    c = http_case
    monkeypatch.setattr(member_scheduler, "async_session_maker", c["maker"])
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            profile = MemberProfile(
                tenant_id=c["school"].id,
                org_unit_id=c["branch"].id,
                name="合成提醒对象",
                org_name=c["branch"].name,
                current_stage="activist",
                stage_joined_on=date(2025, 1, 1),
            )
            db.add(profile)
            await db.commit()
    assert await member_scheduler.scan_member_reminders() >= 1
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            first = (
                (
                    await db.execute(
                        select(MemberReminder).where(MemberReminder.profile_id == profile.id)
                    )
                )
                .scalars()
                .all()
            )
    assert first
    assert await member_scheduler.scan_member_reminders() >= 1
    async with c["maker"]() as db:
        with bypass_tenant_filter():
            second = (
                (
                    await db.execute(
                        select(MemberReminder).where(MemberReminder.profile_id == profile.id)
                    )
                )
                .scalars()
                .all()
            )
    assert {row.id for row in second} == {row.id for row in first}
    assert all(row.tenant_id == c["school"].id for row in second)
    foreign = await call(c, c["other_admin"], "GET", "/member-full/reminders")
    assert foreign["items"] == []
    check = await call(
        c,
        c["author"],
        "POST",
        "/member-full/archive-check",
        {"profile_id": profile.id, "materials_note": {"入党申请书": "合成内容，无签字证据"}},
    )
    assert check["result"]["missing"]
    await call(
        c,
        c["other_admin"],
        "POST",
        "/member-full/archive-check",
        {"profile_id": profile.id, "materials_note": {}},
        status=404,
    )
