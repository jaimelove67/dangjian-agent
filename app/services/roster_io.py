"""名册 Excel 模板读写（T-DATA-06）。

只负责"Excel 文件 ↔ 行数据"与**逐行语法校验**；组织权限、重复判定、落库与审计留在
API/服务层——不把租户与权限判断藏进文件处理代码里。

验收口径（任务清单 T-DATA-06）：可按模板导入 / 导出名册，导入结果可核对，不依赖纯人工
录入；范围限定为文件模板，不承诺全国党员管理信息系统的系统级 API 对接。

决策边界（FR-MEM-10 / T-SEC-13）：导入只做**台账登记**，不代替阶段流转。文件中某人的
阶段与台账不一致时按错误逐行报出，要求走人工流转接口（含材料与时限校验），不静默改写。
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.rules.data_sensitivity import strip_third_party_fields
from app.rules.member_stages import STAGE_LABELS, MemberStage

TEMPLATE_SHEET = "名册"
GUIDE_SHEET = "填写说明"
HEADERS: tuple[str, ...] = (
    "姓名",
    "组织",
    "当前阶段",
    "进入阶段日期",
    "年度",
    "批次",
    "待补材料数",
)
MAX_IMPORT_ROWS = 2000
MAX_EXPORT_ROWS = 5000
_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d")

_VALUE_BY_LABEL = {label: stage.value for stage, label in STAGE_LABELS.items()}
_VALID_STAGES = {stage.value for stage in MemberStage}
_STAGE_CHOICES = "、".join(f"{label}（{stage.value}）" for stage, label in STAGE_LABELS.items())


@dataclass
class RosterRow:
    """一行名册数据；`error` 非空表示该行不能入库。"""

    row_no: int
    name: str = ""
    org_ref: str = ""
    stage: str = ""
    stage_input: str = ""
    stage_joined_on: Optional[date] = None
    year: Optional[int] = None
    batch_no: Optional[str] = None
    pending: int = 0
    error: str = ""


@dataclass
class RosterSheet:
    rows: list[RosterRow] = field(default_factory=list)
    data_rows: int = 0

    @property
    def ok_rows(self) -> list[RosterRow]:
        return [row for row in self.rows if not row.error]


class RosterFormatError(ValueError):
    """整份文件不可用（缺列 / 超行数 / 空表），提示必须能指导用户重做。"""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_stage(value: Any) -> tuple[str, str]:
    raw = _text(value)
    if not raw:
        return "", "当前阶段不能为空"
    if raw in _VALID_STAGES:
        return raw, ""
    if raw in _VALUE_BY_LABEL:
        return _VALUE_BY_LABEL[raw], ""
    return "", f"阶段「{raw}」无效，只能填：{_STAGE_CHOICES}"


def _parse_date(value: Any) -> tuple[Optional[date], str]:
    if value is None or _text(value) == "":
        return None, ""
    parsed: Optional[date]
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        raw = _text(value)
        found: Optional[date] = None
        for fmt in _DATE_FORMATS:
            try:
                found = datetime.strptime(raw, fmt).date()
                break
            except ValueError:
                continue
        if found is None:
            return None, f"进入阶段日期的格式应为 2026-01-31（当前填写：{raw}）"
        parsed = found
    if parsed > date.today():
        return None, f"进入阶段日期 {parsed.isoformat()} 晚于今天，请核对"
    return parsed, ""


def _parse_year(value: Any) -> tuple[Optional[int], str]:
    if value is None or _text(value) == "":
        return None, ""
    raw = _text(value)
    if not raw.isdigit():
        return None, f"年度的格式应为 4 位年份（当前填写：{raw}）"
    year = int(raw)
    if not 2000 <= year <= 2100:
        return None, f"年度 {year} 超出 2000–2100 的范围"
    return year, ""


def _parse_pending(value: Any) -> tuple[int, str]:
    if value is None or _text(value) == "":
        return 0, ""
    raw = _text(value)
    if not raw.isdigit():
        return 0, f"待补材料数应为非负整数（当前填写：{raw}）"
    return int(raw), ""


def read_roster(data: bytes) -> RosterSheet:
    """读入名册模板；行内问题写进 `error`，整份文件不可用时抛 `RosterFormatError`。"""
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl 对损坏文件抛多种异常
        raise RosterFormatError(
            "这个文件不是可读取的 Excel（.xlsx）：请用模板另存/另存为 .xlsx 后重试，"
            "不要把 .xls 或改了后缀的表格直接上传。"
        ) from exc

    sheet_name = TEMPLATE_SHEET if TEMPLATE_SHEET in workbook.sheetnames else workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise RosterFormatError("表格是空的，请下载最新模板后按列填写。") from None

    index: dict[str, int] = {}
    for position, cell in enumerate(header):
        label = _text(cell)
        if label and label not in index:
            index[label] = position
    missing = [name for name in ("姓名", "组织", "当前阶段") if name not in index]
    if missing:
        raise RosterFormatError(
            f"模板缺少列：{'、'.join(missing)}。请下载最新模板，" "只填内容、不要修改表头文字。"
        )

    def take(row: Iterable[Any], column: str) -> Any:
        cells = list(row)
        position = index.get(column)
        if position is None or position >= len(cells):
            return None
        return cells[position]

    parsed_rows: list[RosterRow] = []
    data_rows = 0
    for offset, raw_row in enumerate(rows_iter, start=2):
        if raw_row is None or all(_text(cell) == "" for cell in raw_row):
            continue
        data_rows += 1
        if data_rows > MAX_IMPORT_ROWS:
            raise RosterFormatError(
                f"一次最多导入 {MAX_IMPORT_ROWS} 行，当前文件超过该行数：请拆分后分批导入，"
                "避免一次提交过大而全部失败。"
            )
        row = RosterRow(row_no=offset)
        problems: list[str] = []

        row.name = _text(take(raw_row, "姓名"))
        if not row.name:
            problems.append("姓名不能为空")
        elif len(row.name) > 50:
            problems.append("姓名不能超过 50 字")

        row.org_ref = _text(take(raw_row, "组织"))
        if not row.org_ref:
            problems.append("组织不能为空（填写组织名称或组织编号）")
        elif len(row.org_ref) > 200:
            problems.append("组织名称不能超过 200 字")

        row.stage_input = _text(take(raw_row, "当前阶段"))
        row.stage, stage_error = _parse_stage(row.stage_input)
        if stage_error:
            problems.append(stage_error)

        row.stage_joined_on, date_error = _parse_date(take(raw_row, "进入阶段日期"))
        if date_error:
            problems.append(date_error)

        row.year, year_error = _parse_year(take(raw_row, "年度"))
        if year_error:
            problems.append(year_error)

        batch = _text(take(raw_row, "批次"))
        if len(batch) > 20:
            problems.append("批次不能超过 20 字")
        else:
            row.batch_no = batch or None

        row.pending, pending_error = _parse_pending(take(raw_row, "待补材料数"))
        if pending_error:
            problems.append(pending_error)

        row.error = "；".join(problems)
        parsed_rows.append(row)

    workbook.close()
    return RosterSheet(rows=parsed_rows, data_rows=data_rows)


# ---------------------------------------------------------------- 生成 Excel


def _write_header(sheet: Any, headers: tuple[str, ...], width: tuple[int, ...]) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for position, name in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=position, value=name)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(position)].width = width[position - 1]
    sheet.freeze_panes = "A2"


def _style_header(sheet: Any, width: tuple[int, ...]) -> None:
    _write_header(sheet, HEADERS, width)


def _write_rows(
    sheet: Any, rows: list[list[Any]], *, text_columns: tuple[int, ...] = (4, 5)
) -> None:
    for row_no, values in enumerate(rows, start=2):
        for position, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_no, column=position, value=value)
            if position in text_columns:
                # 日期与年度按文本写入：避免 Excel 按本机区域格式改写用户填的内容
                cell.number_format = "@"


def build_template(org_names: list[str]) -> bytes:
    """下载用的空白模板：表头 + 填写说明（含可用阶段与可用组织）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    _style_header(sheet, (16, 32, 16, 18, 10, 12, 12))

    guide = workbook.create_sheet(GUIDE_SHEET)
    guide.column_dimensions["A"].width = 120
    lines = [
        "填写说明（导入只登记台账，不改动已有人员的当前阶段）",
        "",
        "1. 只在「名册」工作表里填写数据，不要修改第一行表头文字，也不要增删工作表。",
        "2. 姓名：必填，最多 50 字。",
        "3. 组织：必填，填组织名称或组织编号；只能填你权限范围内的组织（见下方清单）。",
        "4. 当前阶段：必填，取值：" + _STAGE_CHOICES,
        f"5. 进入阶段日期：选填，格式 2026-01-31，不能晚于今天（例：{(date.today()).isoformat()}）。",
        "6. 年度 / 批次：选填，年度为 4 位年份、批次最多 20 字；用于跨年度追溯。",
        "7. 待补材料数：选填，非负整数。",
        "",
        "导入结果如何核对：导入后逐行给出「已登记 / 已存在 / 不能导入（原因）」；",
        "任何一行报错都不影响其他行，改正后可以只把这行重新导入。",
        "",
        "关于阶段不一致：如果某人在表里的阶段与台账里的当前阶段不同，该行会报错并说明",
        "原因。阶段变更必须走人工流转（含材料与时限校验），导入不会代替这一步。",
        "",
        "可用组织名称（也可填组织编号）：",
    ]
    lines.extend(f"　- {name}" for name in (org_names or ["（当前权限范围内没有可登记的组织）"]))
    for row_no, text in enumerate(lines, start=1):
        cell = guide.cell(row=row_no, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_no == 1:
            cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_roster(rows: list[dict[str, Any]], *, scope_note: str = "") -> bytes:
    """导出名册：与模板同列，另附导出说明便于核对。"""
    # 第三人字段（备注 / 说明 / 投票意见等）不进批量导出（T-MEM-12）。
    rows = [strip_third_party_fields(row) for row in rows]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    _style_header(sheet, (16, 32, 16, 18, 10, 12, 12))
    _write_rows(
        sheet,
        [
            [
                row.get("name", ""),
                row.get("org_name", ""),
                (
                    STAGE_LABELS.get(MemberStage(row["stage"]), row["stage"])
                    if row.get("stage")
                    else ""
                ),
                row.get("stage_joined_on") or "",
                row.get("year") or "",
                row.get("batch_no") or "",
                row.get("pending") or 0,
            ]
            for row in rows
        ],
    )

    guide = workbook.create_sheet(GUIDE_SHEET)
    guide.column_dimensions["A"].width = 90
    info = [
        "导出说明",
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"导出条数：{len(rows)}",
        f"数据范围：{scope_note or '当前账号权限范围内的组织'}",
        "本表为台账登记用的数据副本，不含原始材料文件；材料仍以系统内记录为准。",
    ]
    for row_no, text in enumerate(info, start=1):
        guide.cell(row=row_no, column=1, value=text).alignment = Alignment(wrap_text=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


UPPER_REPORT_HEADERS: tuple[str, ...] = (
    "事项",
    "类别",
    "姓名",
    "组织",
    "阶段",
    "上级组织",
    "报送文号",
    "报送日期",
    "随报材料",
    "状态",
    "批复结果",
    "批复日期",
    "批复文号",
    "批复说明",
    "是否超过审批时限",
)
UPPER_GUIDE_SHEET = "导出说明"


def build_upper_report_export(rows: list[dict[str, Any]], *, scope_note: str = "") -> bytes:
    """导出上级审批 / 备案报送台账（T-MEM-11）。

    只导**报送要素与批复结果**：随报材料一列是该事项应报的材料**清单**（与阶段材料要求同一
    来源），不是材料原件；政治审查材料等敏感内容不在本表内，因此这份台账可以按需交接口、
    也可以给上级联络人核对，而不会把敏感材料原文带出去。
    """
    # 第三人字段（备注 / 说明 / 投票意见等）不进批量导出（T-MEM-12）。
    rows = [strip_third_party_fields(row) for row in rows]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报送台账"
    _write_header(
        sheet,
        UPPER_REPORT_HEADERS,
        (26, 8, 12, 26, 12, 26, 16, 12, 46, 14, 14, 12, 18, 30, 18),
    )
    _write_rows(
        sheet,
        [
            [
                row.get("item_label") or row.get("item", ""),
                row.get("kind_label", ""),
                row.get("member_name", ""),
                row.get("org_name", ""),
                row.get("stage_label") or row.get("stage", ""),
                row.get("upper_org", ""),
                row.get("letter_no", ""),
                _iso(row.get("submitted_on")),
                "、".join(str(name) for name in (row.get("materials") or [])),
                row.get("status_label", ""),
                row.get("result_label", ""),
                _iso(row.get("decided_on")),
                row.get("reply_no", ""),
                row.get("decide_note", ""),
                "是" if row.get("overdue") else "",
            ]
            for row in rows
        ],
        text_columns=(8, 12),
    )

    guide = workbook.create_sheet(UPPER_GUIDE_SHEET)
    guide.column_dimensions["A"].width = 96
    info = [
        "导出说明",
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"导出条数：{len(rows)}",
        f"数据范围：{scope_note or '当前账号权限范围内的组织'}",
        "本表为上级审批 / 备案的报送台账：事项、报送要素与上级批复结果。",
        "「随报材料」是该事项应报材料的清单（与阶段流转的材料要求同一来源），不含材料原件。",
        "政治审查材料等敏感内容不在本表内；审批与备案的结论以上级党组织的批复为准。",
        "系统只登记报送与批复结果，不代替上级党组织审批，也不自动改动党员的阶段。",
    ]
    for row_no, text in enumerate(info, start=1):
        cell = guide.cell(row=row_no, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_no == 1:
            cell.font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _iso(value: Any) -> str:
    """日期列统一输出 ISO 文本；空值输出空串（导出件不出现 None）。"""
    return value.isoformat() if isinstance(value, date) else ""


def template_headers() -> tuple[str, ...]:
    """供测试与前端文案复用的列定义。"""
    return HEADERS
