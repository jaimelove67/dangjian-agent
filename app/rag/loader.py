"""知识文档解析器与注册机制

对应框架文档 5.2.2："文件解析器按格式注册，新增格式只需注册解析器，不改动入库流程"。

用法::

    registry = build_default_registry()
    parsed = registry.parse(data, file_name="a.pdf")
    parsed.text  # 抽取出的纯文本
"""

from __future__ import annotations

import io
import logging
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

from app.rag.exceptions import ParseError, UnsupportedFormatError

logger = logging.getLogger(__name__)

# 无法直接抽取正文的常见格式（T-WF01-04）
# 系统要能认出它们是什么，并告诉用户该怎么改，而不是丢一句"格式不支持"。
IMAGE_EXTENSIONS: tuple[str, ...] = (
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".gif",
    ".tif",
    ".tiff",
    ".webp",
    ".heic",
)
LEGACY_WORD_EXTENSIONS: tuple[str, ...] = (".doc", ".wps", ".rtf")
TABLE_EXTENSIONS: tuple[str, ...] = (".xls", ".xlsx", ".csv", ".et", ".ppt", ".pptx", ".dps")


@dataclass
class ParsedDocument:
    """解析结果"""

    text: str
    file_name: str = ""
    page_count: Optional[int] = None
    metadata: dict[str, Any] = field(default_factory=dict)


class DocumentParser(ABC):
    """文档解析器基类"""

    name: str = "base"
    extensions: tuple[str, ...] = ()

    @abstractmethod
    def parse(self, data: bytes, *, file_name: str = "") -> ParsedDocument:
        """解析文件字节流"""
        raise NotImplementedError


class PlainTextParser(DocumentParser):
    """纯文本解析器（.txt / .md）"""

    name = "plain_text"
    extensions = (".txt", ".md")

    def parse(self, data: bytes, *, file_name: str = "") -> ParsedDocument:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            # 兼容 GBK 编码的中文文本
            text = data.decode("gbk", errors="ignore")
        return ParsedDocument(text=text, file_name=file_name)


class PdfParser(DocumentParser):
    """PDF 解析器（基础版，基于 pypdf）"""

    name = "pdf"
    extensions = (".pdf",)

    def parse(self, data: bytes, *, file_name: str = "") -> ParsedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - 依赖缺失
            raise ParseError("未安装 pypdf，无法解析 PDF") from exc

        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:
            raise ParseError(f"PDF 解析失败: {exc}") from exc

        return ParsedDocument(text="\n".join(pages), file_name=file_name, page_count=len(pages))


class DocxParser(DocumentParser):
    """Word 解析器（.docx，基于 python-docx）"""

    name = "docx"
    extensions = (".docx",)

    def parse(self, data: bytes, *, file_name: str = "") -> ParsedDocument:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - 依赖缺失
            raise ParseError("未安装 python-docx，无法解析 DOCX") from exc

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ParseError(f"DOCX 解析失败: {exc}") from exc

        text = "\n".join(p.text for p in document.paragraphs)
        return ParsedDocument(text=text, file_name=file_name)


def empty_text_reason(parsed: ParsedDocument) -> str:
    """解析结果没有正文时的分因说明（T-WF01-04：解析失败要有明确提示与可重试的出路）。

    扫描件 PDF、损坏文件与空模板的处置方式完全不同，不能用同一句话打发用户：
    扫描件要说明"文件读到了但没有文字层"，并给出把文字转出来的具体路径。
    """
    ext = Path(parsed.file_name).suffix.lower() or "(无扩展名)"
    if ext == ".pdf":
        if parsed.page_count:
            return (
                f"该 PDF 有 {parsed.page_count} 页，但没有文字层，属于扫描件或图片型 PDF。"
                "系统只能读取 PDF 里的文字层，无法识别图片中的文字，因此不能把它切成检索片段。"
                "请任选一种方式后再上传：\n"
                "① 用扫描仪或 WPS/Adobe 的「文字识别（OCR）」重新导出一份带文字层的 PDF；\n"
                "② 把正文另存为 .docx 或 .txt（可复制文本）后上传；\n"
                "③ 原件确实识别不了时，请把正文要点录入一份 .docx 作为正文版本，原件作附件留存。"
            )
        return "该 PDF 没有可读取的页面，文件可能已损坏或只含目录结构，请重新导出后再上传。"
    return (
        f"{ext} 文件里没有可提取的文本（可能只有图片、或是空白模板）。"
        "请把正文转成可复制的文本后上传：带文字层的 PDF、.docx 或 .txt。"
    )


def unsupported_reason(ext: str, *, supported: list[str]) -> str:
    """不支持的格式要说清"这是什么、为什么不行、该怎么办"（T-WF01-04）。"""
    normalized = ext or "(无扩展名)"
    supported_text = " / ".join(supported)
    if ext in IMAGE_EXTENSIONS:
        return (
            f"图片文件（{normalized}）无法直接抽取正文：系统暂不支持图片文字识别（OCR）。"
            "请把图片里的文字转成可复制的文本后再上传：\n"
            "① 用手机或扫描软件的「提取文字」功能导出 .docx / .txt；\n"
            "② 在 WPS/Word 里插入图片后用「图片转文字」；\n"
            "③ 也可由管理员在知识库中新建正文并录入要点，原件作附件留存。"
        )
    if ext in LEGACY_WORD_EXTENSIONS:
        return (
            f"这是旧版 Word/WPS 格式（{normalized}），系统只能读取 .docx。"
            "请在 Word/WPS 中「另存为 .docx」后重新上传，内容与格式不会丢失。"
        )
    if ext in TABLE_EXTENSIONS:
        return (
            f"表格与演示文件（{normalized}）不能作为制度正文入库："
            f"系统只从 {supported_text} 抽取正文。若是制度附件，请转成 PDF 或 .docx 后上传。"
        )
    return f"不支持的文件格式: {normalized}；当前支持 {supported_text}。"


class ParserRegistry:
    """解析器注册表（按扩展名路由）"""

    def __init__(self) -> None:
        self._by_extension: dict[str, DocumentParser] = {}

    def register(self, parser: DocumentParser) -> None:
        """注册解析器"""
        for ext in parser.extensions:
            self._by_extension[ext.lower()] = parser

    def get(self, file_name: str) -> DocumentParser:
        """按文件名获取解析器"""
        ext = Path(file_name).suffix.lower()
        parser = self._by_extension.get(ext)
        if parser is None:
            raise UnsupportedFormatError(
                unsupported_reason(ext, supported=sorted(self._by_extension))
            )
        return parser

    def parse(self, data: bytes, *, file_name: str) -> ParsedDocument:
        """解析文件（自动选择解析器）"""
        return self.get(file_name).parse(data, file_name=file_name)

    @property
    def supported_extensions(self) -> list[str]:
        """已注册的扩展名"""
        return sorted(self._by_extension)


def build_default_registry() -> ParserRegistry:
    """构建默认注册表（txt/md、pdf、docx）"""
    registry = ParserRegistry()
    for parser in (PlainTextParser(), PdfParser(), DocxParser()):
        registry.register(parser)
    return registry


# 全局默认注册表
default_registry = build_default_registry()
