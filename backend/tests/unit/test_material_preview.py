"""本地提取建议：明确依据、歧义留空、受限解析和无外部调用。"""

import io
import zipfile

from docx import Document

from app.services.material_preview import local_text, preview, suggest_fields


def test_explicit_fields_have_evidence_and_no_stage_or_decision():
    fields, evidence = suggest_fields(
        "2026年10月思想汇报\n汇报人：合成党员\n文档日期：2026年10月1日\n出具单位：合成支部\n摘要：本月学习记录"
    )
    assert fields == {
        "title": "2026年10月思想汇报",
        "material_type": "思想汇报",
        "category": "thought_report",
        "subject_name": "合成党员",
        "doc_date": "2026-10-01",
        "issuer": "合成支部",
        "description": "本月学习记录",
    }
    assert evidence["doc_date"] == "文档日期：2026年10月1日"
    assert set(evidence) == set(fields)


def test_ambiguous_or_implied_values_are_not_inferred():
    fields, _ = suggest_fields(
        "会议讨论了入党申请书和思想汇报。\n某人在2020年1月1日入学。\n姓名：合成甲\n姓名：合成乙\n文档日期：2999-01-01\n材料类型：不支持的类型"
    )
    assert fields == {}
    assert "doc_date" not in suggest_fields("日期：2020-01-01\n日期：2020-02-01")[0]
    assert "doc_date" not in suggest_fields("日期：2020-02-30")[0]


def test_docx_tables_and_gbk_are_read_locally():
    doc = Document()
    doc.add_paragraph("入党申请书")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "申请人：合成人员"
    data = io.BytesIO()
    doc.save(data)
    assert preview(data.getvalue(), "材料.docx")["fields"]["subject_name"] == "合成人员"
    assert (
        preview("思想汇报".encode("gb18030"), "材料.txt")["fields"]["category"] == "thought_report"
    )


def test_images_scanned_pdf_and_corrupt_docx_allow_manual_fallback():
    for name, data in [("材料.png", b"image"), ("材料.pdf", b"broken"), ("材料.docx", b"broken")]:
        result = preview(data, name)
        assert result["status"] == "manual" and not result["fields"] and result["notes"]


def test_docx_expansion_guard_and_text_bound():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * 500000)
    text, notes = local_text(buf.getvalue(), "材料.docx")
    assert not text and "压缩比例" in notes[0]
    text, notes = local_text(("a" * 90000).encode(), "材料.txt")
    assert len(text) == 80000 and notes


def test_native_pdf_text_and_page_cap():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    doc = PdfWriter()
    page = doc.add_blank_page(100, 100)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    commands = DecodedStreamObject()
    commands.set_data(b"BT /F1 10 Tf 10 50 Td (Synthetic material) Tj ET")
    page[NameObject("/Contents")] = commands
    buf = io.BytesIO()
    doc.write(buf)
    assert "Synthetic material" in local_text(buf.getvalue(), "正文.pdf")[0]
    doc.close()
    doc = PdfWriter()
    for _ in range(41):
        doc.add_blank_page(100, 100)
    buf = io.BytesIO()
    doc.write(buf)
    doc.close()
    result = preview(buf.getvalue(), "空白.pdf")
    assert result["status"] == "manual"
    assert any("前 40 页" in note for note in result["notes"])
