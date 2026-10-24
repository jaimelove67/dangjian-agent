"""名册 Excel 模板读写与逐行校验（T-DATA-06）。

纯函数测试：不连库、不调模型，只验证"Excel ↔ 行数据"与报错话术是否可执行。
"""

import io
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from app.services.roster_io import (
    GUIDE_SHEET,
    HEADERS,
    MAX_IMPORT_ROWS,
    TEMPLATE_SHEET,
    RosterFormatError,
    build_roster,
    build_template,
    read_roster,
)


def workbook_bytes(rows, headers=HEADERS) -> bytes:
    """按模板列顺序造一份"用户填好的"Excel。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TEMPLATE_SHEET
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def row(name="张三", org="组织-branch-a", stage="入党积极分子", joined="2024-05-01", **extra):
    values = [name, org, stage, joined, extra.get("year", ""), extra.get("batch", ""), ""]
    return values


class TestTemplate:
    def test_blank_template_has_headers_and_zero_rows(self):
        data = build_template(["第一党支部", "第二党支部"])
        sheet = read_roster(data)
        assert sheet.data_rows == 0
        workbook = load_workbook(io.BytesIO(data))
        assert [cell.value for cell in workbook[TEMPLATE_SHEET][1]] == list(HEADERS)
        guide = "\n".join(
            str(cell.value) for row in workbook[GUIDE_SHEET].iter_rows() for cell in row if cell.value
        )
        # 说明页要能独立指导填写：阶段取值、日期格式、可用组织都在
        assert "入党积极分子（activist）" in guide
        assert "2026-01-31" in guide
        assert "第一党支部" in guide
        # 示例数据不得直接写进数据页，否则用户忘记删除就会被真的导入
        assert workbook[TEMPLATE_SHEET].max_row == 1

    def test_template_without_orgs_still_downloads(self):
        guide = build_template([])
        assert read_roster(guide).data_rows == 0


class TestReadRoster:
    def test_reads_labels_and_values_dates_and_numbers(self):
        data = workbook_bytes(
            [
                row(),
                row(name="李四", stage="probationary", joined="2025/3/9", year=2025, batch="2025-1"),
            ]
        )
        sheet = read_roster(data)
        assert sheet.data_rows == 2
        first, second = sheet.rows
        assert first.stage == "activist" and first.stage_joined_on == date(2024, 5, 1)
        assert second.stage == "probationary" and second.stage_joined_on == date(2025, 3, 9)
        assert second.year == 2025 and second.batch_no == "2025-1"
        assert not first.error and not second.error

    def test_blank_rows_are_skipped(self):
        data = workbook_bytes([row(), ["", "", "", "", "", "", ""], row(name="王五")])
        sheet = read_roster(data)
        assert sheet.data_rows == 2
        assert [item.name for item in sheet.rows] == ["张三", "王五"]

    def test_column_order_and_optional_columns_are_tolerated(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = TEMPLATE_SHEET
        # 用户把列顺序调换、且去掉了选填列：仍应按表头文字正确取值
        sheet.append(["当前阶段", "姓名", "组织", "进入阶段日期"])
        sheet.append(["probationary", "李四", "组织-branch-a", "2025-03-09"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        sheet_result = read_roster(buffer.getvalue())
        parsed = sheet_result.rows[0]
        assert parsed.error == ""
        assert parsed.name == "李四" and parsed.org_ref == "组织-branch-a"
        assert parsed.stage == "probationary" and parsed.stage_joined_on == date(2025, 3, 9)
        assert parsed.year is None and parsed.batch_no is None and parsed.pending == 0

    @pytest.mark.parametrize(
        "values,expected",
        [
            (["", "组织-branch-a", "积极分子", "2024-05-01", "", "", ""], "姓名不能为空"),
            (
                ["张三", "", "积极分子", "2024-05-01", "", "", ""],
                "组织不能为空",
            ),
            (["张三", "组织-a", "未知阶段", "2024-05-01", "", "", ""], "阶段「未知阶段」无效"),
            (
                ["张三", "组织-a", "积极分子", "去年", "", "", ""],
                "进入阶段日期的格式应为 2026-01-31",
            ),
            (
                ["张三", "组织-a", "积极分子", "2999-01-01", "", "", ""],
                "晚于今天",
            ),
            (["张三", "组织-a", "积极分子", "", "民国", "", ""], "年度的格式应为 4 位年份"),
            (["张三", "组织-a", "积极分子", "", "1999", "", ""], "超出 2000–2100"),
            (["张三", "组织-a", "积极分子", "", "", "", "三人"], "待补材料数应为非负整数"),
        ],
    )
    def test_row_level_errors_are_actionable(self, values, expected):
        sheet = read_roster(workbook_bytes([values]))
        assert sheet.data_rows == 1
        assert expected in sheet.rows[0].error
        assert not sheet.ok_rows

    def test_long_batch_name_is_rejected(self):
        values = ["张三", "组织-a", "积极分子", "", "", "批" * 21, ""]
        assert "批次不能超过 20 字" in read_roster(workbook_bytes([values])).rows[0].error

    @pytest.mark.parametrize("missing", ["姓名", "组织", "当前阶段"])
    def test_missing_header_is_reported_by_name(self, missing):
        headers = tuple(name for name in HEADERS if name != missing)
        with pytest.raises(RosterFormatError) as err:
            read_roster(workbook_bytes([row()], headers=headers))
        assert missing in str(err.value)
        assert "不要修改表头文字" in str(err.value)

    def test_corrupt_file_points_at_xlsx(self):
        with pytest.raises(RosterFormatError) as err:
            read_roster(b"PK\x03\x04 not really a workbook")
        message = str(err.value)
        assert ".xlsx" in message and "另存" in message

    def test_too_many_rows_asks_to_split(self):
        rows = [row(name=f"人员{index}") for index in range(MAX_IMPORT_ROWS + 1)]
        with pytest.raises(RosterFormatError) as err:
            read_roster(workbook_bytes(rows))
        assert "拆分后分批导入" in str(err.value)

    def test_future_date_uses_tomorrow(self):
        tomorrow = (date.today() + timedelta(days=1)).isoformat()
        values = ["张三", "组织-a", "积极分子", tomorrow, "", "", ""]
        assert "晚于今天" in read_roster(workbook_bytes([values])).rows[0].error


class TestBuildRoster:
    def test_export_uses_labels_and_adds_scope_note(self):
        data = build_roster(
            [
                {
                    "name": "张三",
                    "org_name": "组织-branch-a",
                    "stage": "activist",
                    "stage_joined_on": "2024-05-01",
                    "year": 2024,
                    "batch_no": "2024-1",
                    "pending": 1,
                }
            ],
            scope_note="全组织范围",
        )
        workbook = load_workbook(io.BytesIO(data))
        assert [cell.value for cell in workbook[TEMPLATE_SHEET][1]] == list(HEADERS)
        exported = [cell.value for cell in workbook[TEMPLATE_SHEET][2]]
        # 导出的阶段必须是中文标签，用户能直接对照，不必查英文枚举
        assert exported[:4] == ["张三", "组织-branch-a", "入党积极分子", "2024-05-01"]
        guide = "\n".join(
            str(cell.value) for row in workbook[GUIDE_SHEET].iter_rows() for cell in row if cell.value
        )
        assert "导出条数：1" in guide and "全组织范围" in guide

    def test_third_party_fields_never_reach_the_export(self):
        """T-MEM-12：即使上游把第三人字段塞进行数据，导出也不落进表格。"""
        data = build_roster(
            [
                {
                    "name": "张三",
                    "org_name": "组织-branch-a",
                    "stage": "activist",
                    "stage_joined_on": "2024-05-01",
                    "year": 2024,
                    "batch_no": "2024-1",
                    "pending": 0,
                    "note": "经向其舅舅王某某核实",
                    "description": "第三人说明",
                    "comment": "投票意见",
                    "peer_review": "党员互评原文",
                }
            ]
        )
        workbook = load_workbook(io.BytesIO(data))
        dumped = "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for marker in ("王某某", "第三人说明", "投票意见", "党员互评原文"):
            assert marker not in dumped

    def test_exported_roster_can_be_read_back_as_import_template(self):
        data = build_roster(
            [
                {
                    "name": "张三",
                    "org_name": "组织-branch-a",
                    "stage": "activist",
                    "stage_joined_on": "2024-05-01",
                    "year": None,
                    "batch_no": None,
                    "pending": 0,
                }
            ]
        )
        sheet = read_roster(data)
        assert sheet.data_rows == 1
        assert sheet.rows[0].error == ""
        assert sheet.rows[0].stage == "activist"
