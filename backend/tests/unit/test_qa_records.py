"""问答记录集中查看与导出（T-SYS-06）。

纯函数测试：查询构造只断言生成的 SQL 条件，Excel 导出只断言表头、行与说明，
都不连库、不调模型。权限与租户边界在接口层，另有 HTTP 集成用例。
"""

import io
from datetime import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.qa_record_io import (
    GUIDE_SHEET,
    HEADERS,
    RECORD_SHEET,
    build_qa_records,
    level_label,
)
from app.core.timeutil import local_time
from app.services.qa_record_service import (
    UNKNOWN_ASKER,
    QaRecordFilters,
    QaRecordRow,
    build_qa_record_query,
    versions_text,
)


def where_of(filters: QaRecordFilters) -> str:
    """只看 WHERE 子句：SELECT 列表里也有列名，整条 SQL 断言会被"顺带"满足。"""
    return str(build_qa_record_query(filters)).split("WHERE", 1)[1]


class TestQueryBuilder:
    def test_tenant_scope_and_soft_deleted_excluded_by_default(self):
        filters = QaRecordFilters(tenant_id="t1")
        where = where_of(filters)
        assert "qa_sessions.tenant_id" in where
        assert "is_deleted IS false" in where
        assert "LEFT OUTER JOIN users" in str(build_qa_record_query(filters))
        assert "ORDER BY qa_sessions.created_at DESC" in str(build_qa_record_query(filters))

    def test_include_deleted_removes_the_soft_delete_condition(self):
        assert "is_deleted IS false" not in where_of(
            QaRecordFilters(tenant_id="t1", include_deleted=True)
        )

    def test_every_filter_reaches_the_query(self):
        where = where_of(
            QaRecordFilters(
                tenant_id="t1",
                user_id="u1",
                keyword="转正",
                data_level="internal",
                refused=True,
                has_evidence=False,
                start_at=datetime(2026, 1, 1, 8, 0),
                end_at=datetime(2026, 2, 1, 8, 0),
            )
        )
        assert "qa_sessions.user_id" in where
        assert "lower(qa_sessions.question) LIKE lower(" in where
        assert "lower(qa_sessions.answer) LIKE lower(" in where
        assert "qa_sessions.data_level" in where
        assert "qa_sessions.refused" in where
        assert "qa_sessions.has_sufficient_evidence" in where
        assert "qa_sessions.created_at >=" in where
        assert "qa_sessions.created_at <=" in where

    def test_absent_filters_do_not_add_conditions(self):
        where = where_of(QaRecordFilters(tenant_id="t1"))
        for column in ("user_id", "data_level", "refused", "has_sufficient_evidence"):
            assert f"qa_sessions.{column} " not in where


class TestScopeNote:
    def test_default_states_deleted_records_are_excluded(self):
        assert QaRecordFilters(tenant_id="t1").scope_note() == "不含本人已删除的记录"

    def test_note_lists_every_active_filter(self):
        note = QaRecordFilters(
            tenant_id="t1",
            user_id="u1",
            keyword="转正",
            data_level="internal",
            refused=True,
            has_evidence=False,
            start_at=datetime(2026, 1, 1, 8, 0),
            end_at=datetime(2026, 2, 1, 8, 0),
        ).scope_note()
        assert "不含本人已删除的记录" in note
        assert "仅提问人 u1" in note
        assert "转正" in note
        assert "数据级别 internal" in note
        assert "仅无依据拒答" in note
        assert "仅无依据" in note
        assert "起 2026-01-01 16:00" in note
        assert "止 2026-02-01 16:00" in note

    def test_note_marks_deleted_records_when_included(self):
        note = QaRecordFilters(tenant_id="t1", include_deleted=True).scope_note()
        assert "含本人已删除的记录" in note


class TestRowMapping:
    def test_missing_user_falls_back_instead_of_blank(self):
        record = SimpleNamespace(
            id="r1",
            user_id="u1",
            question="转正需要什么材料？",
            answer="需要……",
            data_level="internal",
            refused=False,
            has_sufficient_evidence=True,
            retrieved_count=3,
            used_count=2,
            citations=[{"index": 1}],
            warnings=[],
            is_deleted=False,
            created_at=datetime(2026, 1, 1),
            session_id="s1",
            versions=None,
        )
        # 外连接查不到账号（用户已清理）→ 两个字段都是 NULL，也必须给出可读的提问人
        assert QaRecordRow.from_row((record, None, None)).asker == UNKNOWN_ASKER
        # 有账号但没有姓名 → 回落为登录账号，不留空白单元格
        assert QaRecordRow.from_row((record, "", "a1")).asker == "a1"
        # 版本留痕上线前的历史记录（列为 NULL）→ 空字典，不是 None，界面据空判断"未留痕"
        assert QaRecordRow.from_row((record, "", "a1")).versions == {}

    def test_counts_come_from_stored_lists(self):
        record = SimpleNamespace(
            id="r2",
            user_id="u2",
            question="q",
            answer="a",
            data_level="public",
            refused=True,
            has_sufficient_evidence=False,
            retrieved_count=0,
            used_count=0,
            citations=[{"index": 1}, {"index": 2}],
            warnings=["引用[9]不存在，已从答案中移除"],
            is_deleted=True,
            created_at=datetime(2026, 1, 2),
            session_id=None,
            versions={"prompt": {"qa": "v2"}, "workflow": {"WF-03": "wf-03/1"}},
        )
        row = QaRecordRow.from_row((record, "张三", "zhangsan"))
        assert (row.citation_count, row.warning_count) == (2, 1)
        assert row.is_deleted is True
        assert row.refused is True
        # 版本指纹要逐字带出来（SEC-11）：列表与导出都靠它回答"这条答案用了哪个版本"
        assert row.versions["prompt"]["qa"] == "v2"
        assert versions_text(row.versions) == "提示词：问答提示词 v2；工作流：WF-03 wf-03/1"


def record_row(**overrides) -> QaRecordRow:
    base = {
        "id": "r1",
        "user_id": "u1",
        "asker": "张三",
        "username": "zhangsan",
        "question": "转正需要什么材料？",
        "answer": "需要提交转正申请。",
        "data_level": "internal",
        "refused": False,
        "has_sufficient_evidence": True,
        "retrieved_count": 3,
        "used_count": 2,
        "citation_count": 1,
        "warning_count": 0,
        "is_deleted": False,
        "created_at": datetime(2026, 1, 1, 0, 30),
        "session_id": "s1",
    }
    base.update(overrides)
    return QaRecordRow(**base)


class TestLevelLabel:
    def test_known_levels_are_translated(self):
        assert (level_label("public"), level_label("internal")) == ("公开", "内部")

    def test_unknown_level_is_kept_visible(self):
        assert level_label("unexpected") == "unexpected"


class TestLocalTime:
    def test_utc_is_rendered_as_beijing_time(self):
        assert local_time(datetime(2026, 1, 1, 0, 30)) == "2026-01-01 08:30"

    def test_missing_time_is_empty_not_invalid(self):
        assert local_time(None) == ""


class TestExportWorkbook:
    def workbook(self, rows, **kwargs):
        data = build_qa_records(rows, **kwargs)
        return load_workbook(io.BytesIO(data))

    def test_workbook_has_records_and_guide_sheets(self):
        workbook = self.workbook([record_row()])
        assert workbook.sheetnames == [RECORD_SHEET, GUIDE_SHEET]

    def test_headers_match_the_declared_columns(self):
        sheet = self.workbook([])[RECORD_SHEET]
        assert [cell.value for cell in sheet[1]] == list(HEADERS)

    def test_row_values_are_readable(self):
        sheet = self.workbook([record_row()])[RECORD_SHEET]
        values = [cell.value for cell in sheet[2]]
        assert values[0] == "2026-01-01 08:30"
        assert values[1] == "张三"
        assert values[3] == "内部"
        assert values[4] == "已作答"
        assert values[5] == "有依据"
        assert values[6] == "正常"
        # 版本留痕（SEC-11）：这条测试数据没有版本指纹 → 明确写"未留痕"，不假装知道
        assert values[7] == "未留痕"
        assert (values[8], values[9], values[10]) == (3, 2, 1)
        assert values[11] == "转正需要什么材料？"

    def test_version_fingerprint_is_exported(self):
        sheet = self.workbook(
            [record_row(versions={"prompt": {"qa": "v2"}, "workflow": {"WF-03": "wf-03/1"}})]
        )[RECORD_SHEET]
        values = [cell.value for cell in sheet[2]]
        assert values[7] == "提示词：问答提示词 v2；工作流：WF-03 wf-03/1"

    def test_refused_and_deleted_rows_are_labelled(self):
        sheet = self.workbook(
            [record_row(refused=True, has_sufficient_evidence=False, is_deleted=True)]
        )[RECORD_SHEET]
        values = [cell.value for cell in sheet[2]]
        assert values[4] == "无依据拒答"
        assert values[5] == "无依据"
        assert values[6] == "提问人已删除"

    def test_guide_records_scope_count_and_exporter(self):
        guide = self.workbook(
            [record_row()],
            scope_note="不含本人已删除的记录；仅无依据拒答",
            exported_by="李老师（lilaoshi）",
        )[GUIDE_SHEET]
        text = "\n".join(str(cell.value) for cell in guide["A"] if cell.value)
        assert "导出说明" in text
        assert "导出人：李老师（lilaoshi）" in text
        assert "导出条数：1" in text
        assert "筛选口径：不含本人已删除的记录；仅无依据拒答" in text
        assert "北京时间" in text

    def test_empty_result_still_exports_a_usable_file(self):
        workbook = self.workbook([], scope_note="不含本人已删除的记录")
        assert [cell.value for cell in workbook[RECORD_SHEET][1]] == list(HEADERS)
        assert workbook[RECORD_SHEET].max_row == 1
        text = "\n".join(str(cell.value) for cell in workbook[GUIDE_SHEET]["A"] if cell.value)
        assert "导出条数：0" in text
