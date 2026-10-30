"""党员基本信息采集表接口。

- 本人：读取、保存草稿、确认提交；
- 支部组织员：补录组织信息；
- 支部书记、院系管理员、学校管理员：按范围审核；
- 页面与 JSON 不返回完整身份证号，只返回档案中的脱敏快照。
"""

from __future__ import annotations

import io
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.audit import audit_operation
from app.core.desensitization import mask_id_card
from app.core.security import Permission
from app.db.session import get_db
from app.deps import require_permissions
from app.models.member import MemberInfoCollection, MemberProfile
from app.models.user import User, UserRole
from app.schemas.common import APIResponse
from app.schemas.member_info import (
    InfoCollectionList,
    InfoCollectionOut,
    InfoPayload,
    InfoReviewRequest,
    SelfInfoOut,
    SelfInfoPayload,
    SelfInfoProfile,
    SelfInfoSubmit,
)
from app.services.member_info_service import (
    STATUS_LABELS,
    can_edit_branch,
    can_edit_self,
    can_review,
    can_self_fill,
    ensure_scope,
    error_message,
    history_event,
    load_collection,
    merge_payload,
    new_collection,
    now_local,
    own_collection,
    own_profile,
    payload_json,
    person_category,
    refresh_profile_snapshot,
    serialize_collection,
    serialize_summary,
    stage_label,
    suggested_dates,
    validate_complete,
    validate_self,
    visible_org_ids,
)

router = APIRouter()
self_query = require_permissions(Permission.SELF_QUERY)
manage_query = require_permissions(Permission.MEMBER_QUERY)


def _trace(request: Request) -> str | None:
    return getattr(request.state, "trace_id", None)


def _profile_out(
    profile: MemberProfile,
    join_on=None,
    full_on=None,
) -> SelfInfoProfile:
    return SelfInfoProfile(
        id=str(profile.id),
        name=profile.name,
        org_name=profile.org_name,
        stage=profile.current_stage,
        stage_label=stage_label(profile.current_stage),
        id_card_masked=mask_id_card(profile.id_card_no) if profile.id_card_no else "",
        birth_date=profile.birth_date,
        stage_joined_on=profile.stage_joined_on,
        person_category=person_category(profile.current_stage),
        suggested_join_party_on=join_on,
        suggested_full_member_on=full_on,
    )


async def _load_profile(db: AsyncSession, user: User, profile_id: str) -> MemberProfile:
    profile = (
        await db.execute(
            select(MemberProfile).where(
                MemberProfile.id == profile_id,
                MemberProfile.tenant_id == str(user.tenant_id),
                MemberProfile.is_deleted.is_(False),
            )
        )
    ).scalar_one_or_none()
    if profile is None:
        raise HTTPException(404, "党员档案不存在或不在当前租户范围内。")
    return profile


@router.get(
    "/member/info-collection/self",
    response_model=APIResponse[SelfInfoOut],
    summary="本人基本信息采集表",
)
async def own_info_collection(
    request: Request,
    user: User = Depends(self_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[SelfInfoOut]:
    profile = await own_profile(db, user)
    collection = await own_collection(db, user, profile)
    join_on, full_on = await suggested_dates(db, profile)
    return APIResponse(
        data=SelfInfoOut(
            profile=_profile_out(profile, join_on, full_on),
            collection=serialize_collection(collection) if collection else None,
            can_edit=can_self_fill(profile) and can_edit_self(collection),
            can_submit=can_self_fill(profile) and can_edit_self(collection),
        ),
        trace_id=_trace(request),
    )


@router.put(
    "/member/info-collection/self",
    response_model=APIResponse[InfoCollectionOut],
    summary="保存本人基本信息草稿",
)
async def save_own_info_collection(
    body: SelfInfoPayload,
    request: Request,
    user: User = Depends(self_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionOut]:
    profile = await own_profile(db, user)
    if not can_self_fill(profile):
        raise HTTPException(403, "成为预备党员后，才需要填写党员基本信息采集表。")
    collection = await own_collection(db, user, profile)
    if collection is None:
        collection = new_collection(profile, user)
        db.add(collection)
    elif not can_edit_self(collection):
        raise HTTPException(409, "当前采集表已提交，不能直接修改。")

    before = dict(collection.payload or {})
    current = InfoPayload.model_validate(collection.payload or {})
    collection.payload = payload_json(
        merge_payload(current, body, keys=set(SelfInfoPayload.model_fields))
    )
    refresh_profile_snapshot(collection, profile)
    collection.updated_at = now_local()
    await audit_operation(
        db,
        request,
        user,
        action="member_info_save",
        resource_type="member_info_collection",
        resource_id=str(collection.id),
        data_level="sensitive",
        old_value=before,
        new_value=collection.payload,
    )
    await db.commit()
    await db.refresh(collection)
    return APIResponse(
        data=serialize_collection(collection), trace_id=_trace(request)
    )


@router.post(
    "/member/info-collection/self/submit",
    response_model=APIResponse[InfoCollectionOut],
    summary="本人确认提交基本信息采集表",
)
async def submit_own_info_collection(
    body: SelfInfoSubmit,
    request: Request,
    user: User = Depends(self_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionOut]:
    profile = await own_profile(db, user)
    if not can_self_fill(profile):
        raise HTTPException(403, "成为预备党员后，才需要填写党员基本信息采集表。")
    if not body.confirmed:
        raise HTTPException(422, "请先勾选本人确认。")
    collection = await own_collection(db, user, profile)
    if collection is None:
        collection = new_collection(profile, user)
        db.add(collection)
    elif not can_edit_self(collection):
        raise HTTPException(409, "当前采集表已提交，不能重复提交。")

    current = InfoPayload.model_validate(collection.payload or {})
    merged = merge_payload(current, body, keys=set(SelfInfoPayload.model_fields))
    errors = validate_self(merged)
    if errors:
        raise HTTPException(422, error_message(errors))

    at = now_local()
    before = dict(collection.payload or {})
    collection.payload = payload_json(merged)
    refresh_profile_snapshot(collection, profile)
    collection.status = "submitted"
    collection.version = int(collection.version or 0) + 1
    collection.submitted_at = at
    collection.member_confirmed_at = at
    collection.updated_at = at
    collection.history = list(collection.history or []) + [
        history_event(
            "submit",
            user,
            "submitted",
            at,
            "本人确认提交",
            collection.version,
        )
    ]
    await audit_operation(
        db,
        request,
        user,
        action="member_info_submit",
        resource_type="member_info_collection",
        resource_id=str(collection.id),
        data_level="sensitive",
        old_value=before,
        new_value=collection.payload,
    )
    await db.commit()
    await db.refresh(collection)
    return APIResponse(
        data=serialize_collection(collection), trace_id=_trace(request)
    )


def _scope_conditions(user: User, org_ids: set[str]) -> list:
    conditions = [
        MemberInfoCollection.tenant_id == str(user.tenant_id),
        MemberInfoCollection.is_deleted.is_(False),
    ]
    if user.role == UserRole.SCHOOL_ADMIN:
        conditions.append(
            or_(
                MemberInfoCollection.org_unit_id.in_(org_ids),
                MemberInfoCollection.org_unit_id.is_(None),
            )
        )
    else:
        conditions.append(MemberInfoCollection.org_unit_id.in_(org_ids))
    return conditions


@router.get(
    "/member/info-collections",
    response_model=APIResponse[InfoCollectionList],
    summary="党员基本信息采集表列表",
)
async def list_info_collections(
    request: Request,
    status: str | None = Query(None),
    q: str | None = Query(None, max_length=50),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    user: User = Depends(manage_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionList]:
    if user.role == UserRole.SYSTEM_ADMIN:
        raise HTTPException(403, "平台管理员不办理学校党员业务。")
    if status and status not in STATUS_LABELS:
        raise HTTPException(422, "采集表状态无效。")
    org_ids = await visible_org_ids(db, user)
    scope_conditions = _scope_conditions(user, org_ids)
    conditions = list(scope_conditions)
    if status:
        conditions.append(MemberInfoCollection.status == status)
    if q:
        keyword = f"%{q.strip()}%"
        conditions.append(
            or_(
                MemberInfoCollection.member_name.ilike(keyword),
                MemberInfoCollection.org_name.ilike(keyword),
                MemberInfoCollection.member_stage_label.ilike(keyword),
            )
        )

    count_rows = (
        await db.execute(
            select(MemberInfoCollection.status, func.count())
            .where(*scope_conditions)
            .group_by(MemberInfoCollection.status)
        )
    ).all()
    counts = {key: 0 for key in STATUS_LABELS}
    for key, value in count_rows:
        counts[str(key)] = int(value)

    total = int(
        (
            await db.execute(
                select(func.count())
                .select_from(MemberInfoCollection)
                .where(*conditions)
            )
        ).scalar_one()
    )
    rows = (
        await db.execute(
            select(MemberInfoCollection)
            .where(*conditions)
            .order_by(MemberInfoCollection.updated_at.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).scalars().all()
    return APIResponse(
        data=InfoCollectionList(
            items=[serialize_summary(item) for item in rows],
            total=total,
            counts=counts,
        ),
        trace_id=_trace(request),
    )


@router.get(
    "/member/info-collections/{collection_id}",
    response_model=APIResponse[InfoCollectionOut],
    summary="党员基本信息采集表详情",
)
async def get_info_collection(
    collection_id: str,
    request: Request,
    user: User = Depends(manage_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionOut]:
    collection = await load_collection(db, user, collection_id)
    return APIResponse(
        data=serialize_collection(collection), trace_id=_trace(request)
    )


@router.patch(
    "/member/info-collections/{collection_id}",
    response_model=APIResponse[InfoCollectionOut],
    summary="支部补录采集表组织信息",
)
async def update_info_collection(
    collection_id: str,
    body: InfoPayload,
    request: Request,
    user: User = Depends(manage_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionOut]:
    collection = await load_collection(db, user, collection_id)
    if not can_edit_branch(user, collection):
        raise HTTPException(403, "当前账号不能补录这份采集表。")
    profile = await _load_profile(db, user, str(collection.profile_id))
    before = dict(collection.payload or {})
    current = InfoPayload.model_validate(collection.payload or {})
    collection.payload = payload_json(merge_payload(current, body))
    refresh_profile_snapshot(collection, profile)
    at = now_local()
    collection.updated_at = at
    collection.history = list(collection.history or []) + [
        history_event("branch_save", user, collection.status, at, "支部补录核对")
    ]
    await audit_operation(
        db,
        request,
        user,
        action="member_info_branch_save",
        resource_type="member_info_collection",
        resource_id=str(collection.id),
        data_level="sensitive",
        old_value=before,
        new_value=collection.payload,
    )
    await db.commit()
    await db.refresh(collection)
    return APIResponse(
        data=serialize_collection(collection), trace_id=_trace(request)
    )


@router.post(
    "/member/info-collections/{collection_id}/review",
    response_model=APIResponse[InfoCollectionOut],
    summary="审核、退回或备案采集表",
)
async def review_info_collection(
    collection_id: str,
    body: InfoReviewRequest,
    request: Request,
    user: User = Depends(manage_query),
    db: AsyncSession = Depends(get_db),
) -> APIResponse[InfoCollectionOut]:
    collection = await load_collection(db, user, collection_id)
    if user.role == UserRole.SYSTEM_ADMIN:
        raise HTTPException(403, "平台管理员不办理学校党员业务。")
    opinion = body.opinion.strip()
    if not opinion:
        raise HTTPException(422, "请填写处理意见。")
    at = now_local()
    before = dict(collection.payload or {})
    profile = await _load_profile(db, user, str(collection.profile_id))
    current = InfoPayload.model_validate(collection.payload or {})
    if body.payload is not None:
        collection.payload = payload_json(merge_payload(current, body.payload))
    current = InfoPayload.model_validate(collection.payload or {})

    if body.decision == "return":
        if not can_review(user, collection):
            raise HTTPException(403, "当前账号不能退回这份采集表。")
        collection.status = "returned"
        collection.returned_at = at
        collection.returned_by_name = user.name
        collection.return_reason = opinion
        # 保留退回前状态供追溯。
        collection.return_stage = (
            "submitted"
            if user.role == UserRole.BRANCH_SECRETARY
            else "branch_approved"
            if user.role == UserRole.DEPARTMENT_ADMIN
            else "department_approved"
        )
        collection.updated_at = at
        collection.history = list(collection.history or []) + [
            history_event(
                "return", user, "returned", at, opinion, collection.version
            )
        ]
        await audit_operation(
            db,
            request,
            user,
            action="member_info_return",
            resource_type="member_info_collection",
            resource_id=str(collection.id),
            data_level="sensitive",
            old_value=before,
            new_value=collection.payload,
        )
        await db.commit()
        await db.refresh(collection)
        return APIResponse(
            data=serialize_collection(collection), trace_id=_trace(request)
        )

    if body.decision == "archive":
        if user.role != UserRole.SCHOOL_ADMIN or collection.status != "department_approved":
            raise HTTPException(403, "只有学校管理员可以对院系复核通过的采集表备案。")
        collection.status = "archived"
        collection.archived_at = at
        collection.archived_by_name = user.name
        collection.updated_at = at
        collection.history = list(collection.history or []) + [
            history_event(
                "archive", user, "archived", at, opinion, collection.version
            )
        ]
        await audit_operation(
            db,
            request,
            user,
            action="member_info_archive",
            resource_type="member_info_collection",
            resource_id=str(collection.id),
            data_level="sensitive",
            old_value=before,
            new_value=collection.payload,
        )
        await db.commit()
        await db.refresh(collection)
        return APIResponse(
            data=serialize_collection(collection), trace_id=_trace(request)
        )

    if not can_review(user, collection):
        raise HTTPException(403, "当前账号不能审核这份采集表。")
    if user.role == UserRole.BRANCH_SECRETARY:
        errors = validate_complete(current, profile)
        if errors:
            raise HTTPException(422, "请先补全支部信息：" + error_message(errors))
        collection.status = "branch_approved"
        collection.branch_reviewed_at = at
        collection.branch_reviewed_by_name = user.name
        collection.branch_opinion = opinion
        collection.collector_name = collection.collector_name or user.name
        action = "branch_approve"
        status = "branch_approved"
    elif user.role == UserRole.DEPARTMENT_ADMIN:
        errors = validate_complete(current, profile)
        if errors:
            raise HTTPException(422, "请先补全采集表：" + error_message(errors))
        collection.status = "department_approved"
        collection.department_reviewed_at = at
        collection.department_reviewed_by_name = user.name
        collection.department_opinion = opinion
        action = "department_approve"
        status = "department_approved"
    else:
        raise HTTPException(403, "当前账号不能审核这份采集表。")

    collection.payload = payload_json(current)
    refresh_profile_snapshot(collection, profile)
    collection.updated_at = at
    collection.history = list(collection.history or []) + [
        history_event(action, user, status, at, opinion, collection.version)
    ]
    await audit_operation(
        db,
        request,
        user,
        action=f"member_info_{action}",
        resource_type="member_info_collection",
        resource_id=str(collection.id),
        data_level="sensitive",
        old_value=before,
        new_value=collection.payload,
    )
    await db.commit()
    await db.refresh(collection)
    return APIResponse(
        data=serialize_collection(collection), trace_id=_trace(request)
    )


def _docx_text(value) -> str:
    if value is None or value == "":
        return "未填写"
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


@router.get(
    "/member/info-collections/{collection_id}/export.docx",
    summary="导出党员基本信息采集表 Word",
)
async def export_info_collection(
    collection_id: str,
    request: Request,
    user: User = Depends(manage_query),
    db: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    """导出脱敏打印件；完整身份证号不进入接口响应或导出文件。"""

    collection = await load_collection(db, user, collection_id)
    payload = InfoPayload.model_validate(collection.payload or {})
    profile = await _load_profile(db, user, str(collection.profile_id))
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("党员基本信息采集表")
    run.bold = True
    run.font.size = Pt(18)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    meta = document.add_paragraph("附件 2")
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    rows = [
        ("1. 姓名", profile.name),
        ("2. 性别", {"male": "男", "female": "女"}.get(payload.gender, "未填写")),
        ("3. 民族", payload.ethnicity),
        (
            "4. 公民身份证号",
            mask_id_card(profile.id_card_no) if profile.id_card_no else "档案未登记",
        ),
        ("5. 出生日期", profile.birth_date),
        ("6. 学历代码", payload.education_code),
        ("7. 人员类别", person_category(profile.current_stage)),
        ("8. 所在党支部", profile.org_name),
        ("9. 加入党组织日期", payload.join_party_on),
        ("10. 转为正式党员日期", payload.full_member_on),
        ("11. 工作岗位代码", payload.job_code),
        (
            "12. 联系电话",
            " / ".join(
                part
                for part in [
                    payload.mobile,
                    f"{payload.phone_area}-{payload.phone_number}"
                    if payload.phone_area and payload.phone_number
                    else "",
                ]
                if part
            ),
        ),
        ("13. 家庭住址", payload.address),
        (
            "14. 党籍状态",
            {"normal": "正常", "suspended": "停止党籍"}.get(
                payload.party_status, "未填写"
            ),
        ),
        (
            "15. 是否失联党员",
            f"是，失去联系日期：{payload.lost_contact_on}"
            if payload.lost_contact
            else "否",
        ),
        (
            "16. 是否流动党员",
            "是，流向："
            + " ".join(
                part
                for part in [
                    payload.floating_province,
                    payload.floating_city,
                    payload.floating_county,
                ]
                if part
            )
            if payload.floating
            else "否",
        ),
        (
            "17. 填报类型",
            {"new": "当年新发展党员", "supplement": "往年党员补录"}.get(
                payload.filing_type, "未填写"
            ),
        ),
        (
            "18. 是否转接组织关系",
            f"是，转入支部：{payload.transfer_branch}"
            if payload.relation_transfer
            else "否",
        ),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for index, (label, value) in enumerate(rows):
        table.rows[index].cells[0].text = label
        table.rows[index].cells[1].text = _docx_text(value)

    document.add_paragraph()
    signatures = [
        f"党员（签字）：{'已在线确认 ' + collection.member_confirmed_at.isoformat(sep=' ', timespec='minutes') if collection.member_confirmed_at else ''}",
        f"信息采集员（签字）：{collection.collector_name or ''}",
        f"党支部书记（签字）：{collection.branch_reviewed_by_name or ''}",
        f"党总支复核（党总支公章）：{collection.department_reviewed_by_name or ''}",
    ]
    for text in signatures:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_before = Pt(14)
    document.add_paragraph(
        "注：本表一式两份，由党员本人填写并确认，经党支部初核、党总支复核后，"
        "一份交学校党校备案，一份由党支部留存。"
    ).runs[0].font.size = Pt(9)
    document.add_paragraph(
        "系统打印件；身份证号按脱敏值显示，完整号码请按身份证原件核对填写。"
    ).runs[0].font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    await audit_operation(
        db,
        request,
        user,
        action="member_info_export",
        resource_type="member_info_collection",
        resource_id=str(collection.id),
        data_level="sensitive",
        new_value={"format": "docx", "masked_id": True},
    )
    await db.commit()
    filename = quote(
        f"党员基本信息采集表-{collection.member_name}-第{collection.version}版.docx"
    )
    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
